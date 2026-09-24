"""Warnings reviewed before anything is removed: accept, reject, edit a sheet, the Checker."""

import io

import pytest
from conftest import pile_sheet
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from triton.api import app
from triton.validation import apply_section, import_sheets


def ids(wb, code):
    return [i.id for s in wb.sheets for i in s.issues if i.code == code] + [
        i.id for i in wb.issues if i.code == code
    ]


def with_repeat():
    rows = pile_sheet()
    return rows + [rows[1]]  # the first data row again


def test_repeated_rows_are_kept_until_the_removal_is_accepted():
    raw = import_sheets({"Pile(1)-QP": with_repeat(), "Pile(1)-PT-B-Apron": pile_sheet(scale=2)})
    (dup,) = ids(raw, "duplicate_rows_removed")
    kept = apply_section(raw)
    assert len(kept.elements()["Pile(1)"]["QP"].frame) == 6
    removed = apply_section(raw, decisions={dup: "accept"})
    assert len(removed.elements()["Pile(1)"]["QP"].frame) == 5


def test_a_sheet_with_unreadable_rows_is_used_once_leaving_them_out_is_accepted():
    rows = pile_sheet()
    rows[2] = [rows[2][0], "oops", *rows[2][2:]]
    raw = import_sheets({"Pile(1)-QP": rows, "Pile(1)-PT-B-Apron": pile_sheet(scale=2)})
    (bad,) = ids(raw, "non_numeric")
    assert "QP" not in apply_section(raw).elements()["Pile(1)"]
    accepted = apply_section(raw, decisions={bad: "accept"})
    assert len(accepted.elements()["Pile(1)"]["QP"].frame) == 4


def test_rejecting_a_warning_leaves_its_sheet_out():
    raw = import_sheets({"Pile(1)-QP": pile_sheet(), "Pile(1)-PT-B-Apron": pile_sheet()})
    (same,) = ids(raw, "identical_combinations")
    view = apply_section(raw, decisions={same: "reject"})
    assert len(view.elements()["Pile(1)"]) == 1
    assert any(i.code == "rejected" for s in view.sheets for i in s.issues)


def test_piles_with_far_fewer_points_than_their_like_are_flagged():
    raw = import_sheets(
        {
            "Pile(1)-QP": pile_sheet(nodes=range(1, 21)),
            "Pile(2)-QP": pile_sheet(nodes=range(1, 21)),
            "Pile(3)-QP": pile_sheet(nodes=range(1, 21, 4)),  # same length, a quarter of the points
            "Deck-QP": [],
        }
    )
    flagged = [i for i in raw.issues if i.code == "point_count_differs"]
    assert [i.element for i in flagged] == ["Pile(3)"]
    assert "Pile(3) has 5 points" in flagged[0].message
    # Different sizes are not compared.
    sized = apply_section(raw, sizes={"Pile(1)": "a", "Pile(2)": "a", "Pile(3)": "b"})
    assert not [i for i in sized.issues if i.code == "point_count_differs"]


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("TRITON_DATA_DIR", str(tmp_path))
    return TestClient(app)


def xlsx(sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_see_edit_and_download_a_sheet(client):
    p = client.post("/api/projects", json={"element_names": ["Pile(1)"]}).json()
    url = f"/api/projects/{p['id']}/sections/{p['sections'][0]['id']}"
    rows = pile_sheet()
    rows[3] = [rows[3][0], "n/a", *rows[3][2:]]
    data = xlsx({"Pile(1)-QP": rows, "Pile(1)-PT-B-Apron": pile_sheet(scale=2)})
    assert client.post(f"{url}/workbook", files={"file": ("s.xlsx", data)}).status_code == 200

    page = client.get(f"{url}/workbook/sheet", params={"name": "Pile(1)-QP"}).json()
    assert page["editable"] and page["flagged_rows"] == [4]
    assert page["rows"][4 - 1 - page["start"]][1] == "n/a"

    # Fix the cell in Triton: the sheet is read again and the warning is gone.
    fixed = client.put(
        f"{url}/workbook/sheet",
        params={"name": "Pile(1)-QP"},
        json={"edits": [{"row": 4, "col": 1, "value": "3"}]},
    ).json()
    assert not [i for i in fixed["issues"] if i["code"] == "non_numeric"]
    assert fixed["file"].endswith("(edited)")

    # The Checker: a list sheet linking to the flagged rows, and the flagged sheets.
    rows[2] = [rows[2][0], "bad", *rows[2][2:]]
    client.post(
        f"{url}/workbook",
        files={"file": ("s.xlsx", xlsx({"Pile(1)-QP": rows, "Pile(1)-PT-B-Apron": pile_sheet()}))},
    )
    got = client.get(f"{url}/workbook/checker.xlsx")
    assert got.status_code == 200
    wb = load_workbook(io.BytesIO(got.content))
    assert wb.sheetnames[0] == "Triton checker" and "Pile(1)-QP" in wb.sheetnames
    ws = wb["Pile(1)-QP"]
    assert ws["A3"].comment and "text where a number" in ws["A3"].comment.text
    # Uploaded again, the list sheet is skipped.
    buf = io.BytesIO(got.content)
    again = client.post(
        f"{url}/workbook?mode=update", files={"file": ("checker.xlsx", buf.getvalue())}
    ).json()
    assert "Triton checker" not in [s["name"] for s in again["sheets"]]
