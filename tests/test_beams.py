import math

import numpy as np
import pytest
from conftest import PLATE_HEADER, pile_sheet

from triton.design.beams import Support, design_beam, layout, shear_stations, station_forces
from triton.design.circular import ConcreteLaw, SteelLaw
from triton.design.crack import crack_width, restraint_factor
from triton.design.governing import workbook
from triton.design.rect import Bars, RectSection
from triton.design.runner import run_section
from triton.geometry import section_geometry
from triton.project import BeamInput, DesignSettings, PileInput, Section
from triton.validation import import_sheets

LAWS = (ConcreteLaw(40), SteelLaw(500))


def test_rectangular_capacity_and_cracked_stress_match_hand_calcs():
    # 1000 x 1000, 5Ø25 at d = 940, 2Ø12 on top.
    bars = Bars.join(Bars.row(5, 25, -440, 400), Bars.row(2, 12, 440, 400))
    sec = RectSection(1000, 1000, bars, *LAWS)
    a_s = 5 * math.pi * 25**2 / 4
    fyd, fcd = 500 / 1.15, 0.85 * 40 / 1.5
    x = a_s * fyd / (0.8 * 1000 * fcd)
    assert sec.m_rd("v", np.array([1]), np.array([0.0]))[0] == pytest.approx(
        a_s * fyd * (940 - 0.4 * x) / 1e6, rel=0.01
    )
    # Hogging is weaker: only the two top bars.
    assert (
        sec.m_rd("v", np.array([-1]), np.array([0.0]))[0]
        < 0.2 * sec.m_rd("v", np.array([1]), np.array([0.0]))[0]
    )
    # Cracked elastic section, αe = 200 / (35 / 3): x = k·d, σs = M / (As (d − x/3)).
    e_c = 35_000 / 3
    ae, rho = 200_000 / e_c, a_s / (1000 * 940)
    k = math.sqrt(2 * ae * rho + (ae * rho) ** 2) - ae * rho
    r = sec.cracked(0.0, 500.0, e_c)
    assert r["x"] == pytest.approx(k * 940, rel=0.02)
    assert -r["stress"][:5].min() == pytest.approx(500e6 / (a_s * (940 - k * 940 / 3)), rel=0.01)


def test_biaxial_interaction_to_5_8_9():
    bars = Bars.join(Bars.row(4, 25, -440, 440), Bars.row(4, 25, 440, 440))
    sec = RectSection(1000, 1000, bars, *LAWS)
    mv = sec.m_rd("v", np.array([1]), np.array([0.0]))[0]
    mh = sec.m_rd("h", np.array([1]), np.array([0.0]))[0]
    # At N = 0, a = 1: the two ratios add.
    u = sec.utilisation(np.array([0.0]), np.array([mv / 2]), np.array([mh / 2]))
    assert u[0] == pytest.approx(1.0, abs=0.01)
    assert sec.utilisation(np.array([0.0]), np.array([mv]))[0] == pytest.approx(1.0, abs=0.01)
    # Pure tension beyond the bars' capacity fails.
    assert sec.utilisation(np.array([-1.1 * bars.total * 500 / 1.15 / 1e3]), np.array([0.0]))[0] > 1


def test_crack_width_to_7_3_4():
    # σs = 250 MPa, Ø20 @ 150 at c = 50, h = 1000, d = 940, x = 250, fctm 3.5, Ecm 35 GPa.
    area = 1000 / 150 * math.pi * 100
    c = crack_width(
        250, h=1000, d=940, x=250, b=1000, area=area, phi=20, cover=50, spacing=150, fctm=3.5, ecm=35_000
    )
    hc = min(2.5 * 60, 750 / 3, 500)
    rho = area / (1000 * hc)
    strain = max((250 - 0.4 * 3.5 / rho * (1 + 200 / 35 * rho)) / 200_000, 0.6 * 250 / 200_000)
    sr = 3.4 * 50 + 0.8 * 0.5 * 0.425 * 20 / rho
    assert c["wk"] == pytest.approx(sr * strain, abs=0.002)
    assert (
        crack_width(
            -10, h=1000, d=940, x=250, b=1000, area=area, phi=20, cover=50, spacing=150, fctm=3.5, ecm=35_000
        )["wk"]
        == 0
    )


def test_restraint_factor_from_length_over_height():
    assert restraint_factor(30, 2) == pytest.approx((13 / 16) ** (1 / 15))
    assert restraint_factor(4, 2) == pytest.approx((1 / 12) ** 0.5)
    assert restraint_factor(1, 2) == 0
    assert restraint_factor(60, 2) > restraint_factor(10, 2)


def beam_rows(fn, width=2.0, length=12.0, z=2.7):
    """Plate rows of a beam along Y centred on X = 0; fn(x, y) -> the 8 plate actions."""
    rows = [PLATE_HEADER]
    node = 1
    for y in np.arange(0.0, length + 1e-9, 0.25):
        for x in np.linspace(-width / 2, width / 2, 5):
            vals = fn(x, y)
            row = ["Plate\\_1\\_1", node, 1, float(x), float(y), z]
            for v in vals:
                row += [v, min(v, 0.0), max(v, 0.0)]
            rows.append(row)
            node += 1
    return rows


def uniform(n2=0.0, m22=0.0, q23=0.0, m11=0.0, grad_n2=0.0):
    # N_1, N_2, Q_12, Q_23, Q_13, M_11, M_22, M_12
    return lambda x, y: [0.0, n2 + grad_n2 * x, 0.0, q23, 0.0, m11, m22, 0.0]


def test_section_forces_are_integrated_across_the_width():
    raw = {
        "Front Beam-PT-B-Apron": beam_rows(uniform(n2=-50.0, m22=100.0, q23=30.0, grad_n2=20.0)),
        "Front Beam-QP": beam_rows(uniform(m22=60.0)),
    }
    sheets = import_sheets(raw).elements()["Front Beam"]
    lay = layout(sheets, {"1": "X", "2": "Y"})
    assert (lay.along, lay.span_local, lay.width) == ("Y", "2", 2.0)
    f = station_forces(sheets["PT-B-Apron"].frame, lay, 1.0, np.array([6.0]))
    r = f.iloc[0]
    assert r["N"] == pytest.approx(100.0)  # 50 kN/m tension × 2 m, concrete sign
    assert r["Mv"] == pytest.approx(200.0)
    assert r["V"] == pytest.approx(60.0)
    # N2 = 20·x kN/m across a 2 m width: Mh = 20·2³/12, compression on the −x side.
    assert abs(r["Mh"]) == pytest.approx(20 * 8 / 12)
    assert station_forces(sheets["PT-B-Apron"].frame, lay, -1.0, np.array([6.0])).iloc[0][
        "Mv"
    ] == pytest.approx(-200.0)


def test_shear_falls_back_to_midway_between_close_supports():
    s = np.arange(0, 6.01, 0.1)
    sup = [Support("P", 0.0, 0.0, 0.6), Support("P", 3.0, 0.0, 0.6), Support("P", 6.0, 0.0, 0.6)]
    keep, fallback = shear_stations(s, sup, 1.9)
    assert fallback
    assert sorted(np.round(s[keep], 2)) == [1.5, 4.5]
    keep, fallback = shear_stations(np.arange(0, 12.01, 0.1), sup[:1], 1.0)
    assert not fallback and keep.sum() == len(np.arange(1.6, 12.01, 0.1))


def pile_line(y_positions, z_top=2.7):
    rows = pile_sheet()[:1]
    for i, y in enumerate(y_positions):
        for j, z in enumerate((z_top, z_top - 5)):
            row = ["EmbeddedBeam\\_1\\_1", 100 + 2 * i + j, 1, 0.0, float(y), z]
            row += [-1000.0, -1001.0, 0.0] + [0.0, 0.0, 0.0] * 5 + [1.0, 2.0, 3.0, "N/A"]
            rows.append(row)
    return rows


def test_design_beam_with_supports_and_export():
    peak = uniform(m22=150.0, q23=80.0, m11=300.0)

    def with_peaks(x, y):
        v = peak(x, y)
        if math.hypot(x, y - 6.0) < 0.55:  # FE peak inside the pile at y = 6
            v[6] = 5000.0
        return v

    raw = {
        "Front Beam-PT-B-Apron": beam_rows(with_peaks),
        "Front Beam-QP": beam_rows(uniform(m22=100.0, m11=200.0)),
        "Pile(1)-PT-B-Apron": pile_line([0.0, 6.0, 12.0]),
        "Pile(1)-QP": pile_line([0.0, 6.0, 12.0]),
    }
    wb = import_sheets(raw)
    truss = {"crane_load": 0, "bollard_slab_thickness": None}
    els = {"Front Beam": BeamInput(depth=1500, truss=truss), "Pile(1)": PileInput(head_level=2.7)}
    faces = DesignSettings(beam_support_results="faces")
    d = design_beam(
        "Front Beam",
        els["Front Beam"],
        faces,
        wb.elements()["Front Beam"],
        section_geometry(wb),
        els,
        {"1": "X", "2": "Y"},
    )
    # The QP diagrams: the same positions, the QP envelope and its crack width over the limit.
    assert d["profile_qp"] and {"s", "u", "Mv_max", "Mv_min"} <= set(d["profile_qp"][0])
    assert max(q["Mv_max"] for q in d["profile_qp"]) < max(q["Mv_max"] for q in d["profile"])
    # By default every result is designed, the peak inside the pile too, as the office's beam designs.
    every = design_beam(
        "Front Beam",
        els["Front Beam"],
        DesignSettings(),
        wb.elements()["Front Beam"],
        section_geometry(wb),
        els,
        {"1": "X", "2": "Y"},
    )
    assert every["support_results"] == "all" and every["bending"]["extremes"]["Mv"]["max"] > 5000
    assert [q["s"] for q in every["supports"]] == [0.0, 6.0, 12.0]
    assert [q["s"] for q in d["supports"]] == [0.0, 6.0, 12.0]
    # 6 m between piles on a 1.5 m deep beam: the truss tie, not bending, sets the bottom bars.
    assert d["passed"] and d["utilisation"] <= 1
    assert (
        d["truss"]["utilisation"] > 0.9
        and d["truss"]["As_provided_mm2"] > d["truss"]["cases"][0]["As_req_mm2"]
    )
    # The 5000 kNm/m peak lies inside the pile at y = 6 and is left out.
    assert d["bending"]["extremes"]["Mv"]["max"] < 1000
    assert set(d["cracks"]) == {"bottom"} and d["cracks"]["bottom"]["limit"] == 0.2
    assert d["restraint"]["R"] == pytest.approx(restraint_factor(58, 1.5), abs=1e-3)
    assert d["shear"]["link"]["legs"] >= 2 and d["transverse"]["bottom"]["label"].startswith("Ø")
    assert d["steel"]["kg_per_m3"] > 0
    assert d["truss"]["spacing_m"] == 6.0 and d["truss"]["spacing_from"] == "workbook"
    assert len(d["governing_sets"][0]["uls"]) == 7 and len(d["governing_sets"][0]["qp"]) == 7
    # Crack view: each QP set with its crack at its tension face; bands of wk / limit along the beam.
    qp = d["governing_sets"][0]["qp"]
    assert max(r["crack"]["wk_mm"] for r in qp) == pytest.approx(d["cracks"]["bottom"]["wk"], abs=2e-3)
    assert all(r["crack"]["face"] == ("bottom" if r["M3_kNm"] >= 0 else "top") for r in qp)
    worst = d["cracks"]["bottom"]["wk"] / d["cracks"]["bottom"]["limit"]
    assert max(b[3] for b in d["crack_bands"]) == pytest.approx(worst, abs=2e-3)

    res = run_section(faces, Section(elements=els), wb)
    assert [b["element"] for b in res["beams"]] == ["Front Beam"]
    from io import BytesIO

    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(workbook("P", "S", res)))["Concrete"]
    titles = [r[0] for r in ws.iter_rows(values_only=True) if r[0] and str(r[0]).startswith("Front Beam")]
    assert titles and "M3 vertical" in titles[0]
    from triton.design.export import pile_cages

    (jb,) = pile_cages("P", res, "S")["beams"]
    assert jb["element"] == "Front Beam" and len(jb["bars"]) == len(res["beams"][0]["cage"]["bars"])
    assert jb["links"]["spacing_mm"] > 0 and {"top", "bottom"} <= set(jb["transverse"])


def test_bollard_ties_as_the_office_drawing():
    from triton.design.bollard import check_bollard
    from triton.project import Bollard, TieBars

    d = check_bollard(Bollard(), "C40/50", DesignSettings())
    # 150 t × 0.75 (the report's mooring factor) × g square to the quay; 2Ø32 straight and 3Ø32 at
    # ±45°, 8.11° down.
    assert d["F_Ed_kN"] == round(0.75 * 150 * 9.81)
    assert d["passed"] and d["tie_utilisation"] == pytest.approx(0.511, abs=0.005)
    leading = check_bollard(Bollard(load_factor=1.5), "C40/50", DesignSettings())
    assert leading["tie_utilisation"] == pytest.approx(1.022, abs=0.005)
    fyd = 500 / 1.15
    bar = math.pi * 32**2 / 4 * fyd * math.cos(math.radians(8.11)) / 1e3
    assert d["R_kN"] == pytest.approx(2 * bar + 6 * bar * math.cos(math.pi / 4), abs=2)
    assert d["tie_utilisation"] == pytest.approx(d["F_Ed_kN"] / d["R_kN"], abs=2e-3)
    more = check_bollard(
        Bollard(ties=[TieBars(count=2), TieBars(count=6, angle=45), TieBars(count=6, angle=-45)]),
        "C40/50",
        DesignSettings(),
    )
    assert more["passed"] and more["laps"][0]["l0_mm"] <= 1600


def test_truss_model_as_the_office_report():
    from triton.design.truss import check_truss, spacing_from_supports
    from triton.project import FrontBeamTruss

    # Report 5.3: 4.5 x 2.0 m beam, z = 1.726 m, king piles 3.211 m apart (loads rounded to 3.2 m),
    # 56Ø25 bottom bars (275 cm²).
    d = check_truss(FrontBeamTruss(), 4500, 2000, 1726, 56 * math.pi * 25**2 / 4, 3.211)
    typical, bollard = d["cases"]
    assert d["theta_deg"] == pytest.approx(47.07, abs=0.01)
    assert typical["P_beam_kN"] == pytest.approx(4104, rel=0.004)  # 410 t
    assert typical["P_slab_kN"] == pytest.approx(504, rel=0.004)  # 50.4 t
    assert bollard["P_slab_kN"] == pytest.approx(672, rel=0.004)  # 67.2 t, slab thickened to 1.4 m
    assert d["tie_factor"] == pytest.approx(0.465, abs=0.002)
    assert typical["As_req_mm2"] == pytest.approx(21410, rel=0.005)  # 214.1 cm²
    assert d["utilisation"] == pytest.approx(222 / 275, abs=0.005)  # the report's 0.81
    assert spacing_from_supports([0.0, 3.2, 6.4, 9.6, 12.8]) == pytest.approx(3.2)
    assert check_truss(FrontBeamTruss(), 4500, 2000, 1726, 1.0, None)["utilisation"] is None


def test_torsion_steel_comes_out_of_the_cage():
    def twist(m12):
        return lambda x, y: [0.0, 0.0, 0.0, 80.0, 0.0, 300.0, 400.0, m12]

    def run(m12):
        raw = {"Front Beam-PT-B-Apron": beam_rows(twist(m12)), "Front Beam-QP": beam_rows(twist(0.0))}
        wb = import_sheets(raw)
        el = BeamInput(depth=1500)
        return design_beam(
            "Front Beam",
            el,
            DesignSettings(),
            wb.elements()["Front Beam"],
            [],
            {"Front Beam": el},
            {"1": "X", "2": "Y"},
        )

    plain, twisted = run(0.0), run(600.0)
    t = twisted["bending"]["torsion_steel"]
    assert t["asl_mm2"] > 0
    per = 2 * (2000 + 1500)
    if twisted["cage"]["side"]["count"]:
        assert t["side_mm2"] == pytest.approx(t["asl_mm2"] * 1500 / per, abs=1)
    assert t["top_mm2"] + t["bottom_mm2"] + 2 * t["side_mm2"] == pytest.approx(t["asl_mm2"], abs=2)
    assert "torsion_steel" not in plain["bending"]
    assert twisted["steel"]["longitudinal_kg_per_m"] >= plain["steel"]["longitudinal_kg_per_m"]
    assert any("6.3.2(3)" in n for n in twisted["notes"])


def test_peak_times_width_takes_the_largest_nodal_values():
    raw = {
        "Front Beam-PT-B-Apron": beam_rows(
            lambda x, y: [0.0, -50.0, 0.0, 30.0 + 10 * x, 0.0, 0.0, 100.0 + 50 * x, 0.0]
        )
    }
    sheets = import_sheets(raw).elements()["Front Beam"]
    lay = layout(sheets, {"1": "X", "2": "Y"})
    f = station_forces(sheets["PT-B-Apron"].frame, lay, 1.0, np.array([6.0]), peak_width=4.5)
    assert sorted(f["Mv"]) == pytest.approx([50 * 4.5, 150 * 4.5])
    assert np.allclose(f["V"], 40 * 4.5)
    assert np.allclose(f["N"], 100.0)  # still integrated over the model's 2 m


def test_beam_reports_what_sets_each_face():
    settings = DesignSettings()
    raw = {
        "Front Beam-PT-B-Apron": beam_rows(uniform(m22=100.0, q23=30.0)),
        "Front Beam-QP": beam_rows(uniform(m22=60.0)),
    }
    sheets = import_sheets(raw).elements()["Front Beam"]
    d = design_beam(
        "Front Beam", BeamInput(kind="front_beam", width=2000, depth=1600), settings, sheets, [], {}, None
    )
    faces = {f["face"]: f for f in d["faces"]}
    assert set(faces) == {"top", "bottom", "side"}
    assert faces["top"]["needs_mm2"]["minimum"] > 0 and faces["top"]["governed_by"]
    assert any("peak nodal" in n for n in d["notes"])


def test_bars_set_by_the_user_are_checked_as_they_are():
    from triton.project import BeamCage

    raw = {
        "Front Beam-PT-B-Apron": beam_rows(uniform(m22=400.0, q23=30.0)),
        "Front Beam-QP": beam_rows(uniform(m22=300.0)),
    }
    sheets = import_sheets(raw).elements()["Front Beam"]
    beam = BeamInput(kind="front_beam", width=2000, depth=1600)
    auto = design_beam("Front Beam", beam, DesignSettings(), sheets, [], {}, None)
    light = BeamCage(
        top={"count": 10, "diameter": 16},
        bottom={"count": 10, "diameter": 16},
        side={"count": 4, "diameter": 16},
    )
    d = design_beam("Front Beam", beam, DesignSettings(), sheets, [], {}, None, light)
    assert d["user_set"] and not auto["user_set"]
    assert d["cage"]["bottom"]["count"] == 10 and d["cage"]["bottom"]["phi"] == 16
    assert d["cracks"]["bottom"]["wk"] > auto["cracks"]["bottom"]["wk"]
    assert any("set by you" in n for n in d["notes"])
    # The drawing shows the bars at their real size.
    assert {bar[2] for bar in d["cage"]["bars"]} == {16}
    tight = BeamCage(**{**light.model_dump(), "bottom": {"count": 60, "diameter": 32}})
    t = design_beam("Front Beam", beam, DesignSettings(), sheets, [], {}, None, tight)
    assert not t["passed"] and any("clear spacing" in n for n in t["notes"])


def test_beam_top_and_bottom_crack_limits():
    raw = {
        "Front Beam-PT-B-Apron": beam_rows(uniform(m22=150.0, m11=300.0)),
        "Front Beam-QP": beam_rows(uniform(m22=100.0, m11=200.0)),
        "Pile(1)-PT-B-Apron": pile_line([0.0, 6.0, 12.0]),
        "Pile(1)-QP": pile_line([0.0, 6.0, 12.0]),
    }
    wb = import_sheets(raw)
    truss = {"crane_load": 0, "bollard_slab_thickness": None}
    beam = BeamInput(depth=1500, truss=truss, crack_width_limit=0.3, crack_width_limit_bottom=0.1)
    els = {"Front Beam": beam, "Pile(1)": PileInput(head_level=2.7)}
    d = design_beam(
        "Front Beam",
        beam,
        DesignSettings(),
        wb.elements()["Front Beam"],
        section_geometry(wb),
        els,
        {"1": "X", "2": "Y"},
    )
    limits = {f: c["limit"] for f, c in d["cracks"].items()}
    assert limits.get("bottom", 0.1) == 0.1 and limits.get("top", 0.3) == 0.3
    assert all(c["wk"] <= c["limit"] + 1e-9 for c in d["cracks"].values())
