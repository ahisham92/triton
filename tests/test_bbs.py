"""Bar bending schedule: bar marks, shape codes, cut lengths and weights of the designed bars."""

from __future__ import annotations

import io
import json
import math
from pathlib import Path

from openpyxl import load_workbook

from triton.bbs import _pieces, kg_per_m, schedule, workbook
from triton.project import Project

SAMPLE = json.loads((Path(__file__).parent / "data" / "sample_cages.json").read_text("utf-8"))


def test_long_bars_are_lapped_stock_lengths():
    assert _pieces(9000, 32, 12000, 45) == [(9000, 1)]
    pieces = _pieces(30000, 25, 12000, 45)
    lap = 45 * 25
    assert sum(n for _, n in pieces) == 3
    assert sum(length * n for length, n in pieces) - 2 * lap == 30000


def test_unit_weights_follow_bs_8666():
    assert math.isclose(kg_per_m(32), 6.313, abs_tol=0.002)
    assert math.isclose(kg_per_m(10), 0.617, abs_tol=0.001)


def test_schedule_from_cages(monkeypatch):
    import triton.bbs as b

    monkeypatch.setattr(b, "pile_cages", lambda *a, **k: SAMPLE)
    monkeypatch.setattr(b, "named_parts", lambda r: r)
    results = {
        "piles": [
            {"element": "Pile(1)", "shear": {"zones": [{"top": 2.7, "bottom": -20.0, "spacing_mm": 175}]}}
        ]
    }
    rows = schedule(Project(), results)
    marks = [r["mark"] for r in rows]
    assert len(marks) == len(set(marks))
    assert {"00", "75", "51", "99"} <= {r["shape"] for r in rows}
    ring = next(r for r in rows if r["element"] == "Pile(1)" and r["shape"] == "75")
    assert ring["dims"]["A"] == 1050 and ring["total"] == ring["per_member"] * 8
    assert all(r["length_mm"] % 25 == 0 and r["length_mm"] <= 12000 for r in rows)
    beam = [r for r in rows if r["element"] == "Front Beam" and r["shape"] == "00" and "bars" in r["where"]]
    assert sum(r["total"] for r in beam if r["where"] == "Top bars") % 19 == 0
    wb = load_workbook(io.BytesIO(workbook(Project(), "Section 1", rows)))
    assert wb.sheetnames == ["Schedule", "Weights", "How lengths are taken"]
    assert wb["Schedule"]["A5"].value == rows[0]["mark"]
