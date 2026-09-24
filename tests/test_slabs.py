import math

import numpy as np
import pandas as pd
import pytest
from conftest import PLATE_HEADER, pile_sheet

from triton.design.runner import run_section
from triton.design.slabs import (
    auto_stations,
    bar_options,
    required_as,
    station_text,
    strip_frame,
    wood_armer,
    zones_for,
)
from triton.project import (
    CraneArea,
    DesignSettings,
    PileInput,
    PunchingDepth,
    Section,
    SlabInput,
    SlabMesh,
    SlabStrips,
)
from triton.validation import import_sheets


def test_wood_armer_moments():
    m = wood_armer(np.array([100.0, -80.0]), np.array([50.0, -40.0]), np.array([20.0, 10.0]))
    # All sagging: bottom = M + |Mxy|, nothing on top.
    assert (m["bottom_x"][0], m["bottom_y"][0]) == (120.0, 70.0)
    assert (m["top_x"][0], m["top_y"][0]) == (0.0, 0.0)
    # All hogging: top = M − |Mxy|, nothing at the bottom.
    assert (m["top_x"][1], m["top_y"][1]) == (-90.0, -50.0)
    assert (m["bottom_x"][1], m["bottom_y"][1]) == (0.0, 0.0)


def test_required_steel_per_metre():
    fyd = 500 / 1.15
    a, k, _ = required_as(np.array([500.0]), np.array([0.0]), 1000, 900, 40, fyd)
    assert k[0] == pytest.approx(500e6 / (1000 * 900**2 * 40))
    assert a[0] == pytest.approx(500e6 / (0.95 * 900 * fyd))  # z capped at 0.95d
    # Axial tension adds steel, compression takes some away.
    t, _, _ = required_as(np.array([500.0]), np.array([-200.0]), 1000, 900, 40, fyd)
    c, _, _ = required_as(np.array([500.0]), np.array([200.0]), 1000, 900, 40, fyd)
    assert c[0] < a[0] < t[0]


def test_zones_merge_cells_into_rectangles():
    options = bar_options(DesignSettings())
    cells = pd.DataFrame({"i": [0, 1, 2, 0, 1, 2], "j": [0, 0, 0, 1, 1, 1], "idx": [0, 0, 5, 0, 0, 5]})
    ok = np.zeros((6, len(options)), bool)
    ok[:, 5:] = True
    ok[[0, 1, 3, 4], :] = True
    z = zones_for(cells, ok, options, 1.0, 0.0, 0.0, "X")
    assert z["basic_index"] == 0
    assert z["zones"] == [
        {
            "x": [2.0, 3.0],
            "y": [0.0, 2.0],
            "label": z["zones"][0]["label"],
            "as_mm2_per_m": round(options[5][0]),
        }
    ]


def deck_rows(fn, x_range=(-8.0, 0.0), y_range=(-4.0, 4.0), step=0.5, z=2.7):
    rows = [PLATE_HEADER]
    node = 1
    for x in np.arange(x_range[0], x_range[1] + 1e-9, step):
        for y in np.arange(y_range[0], y_range[1] + 1e-9, step):
            row = ["Plate\\_1\\_1", node, 1, float(x), float(y), z]
            for v in fn(x, y):
                row += [v, min(v, 0.0), max(v, 0.0)]
            rows.append(row)
            node += 1
    return rows


def piles_at(points, n=1500.0, z_top=2.7):
    rows = pile_sheet()[:1]
    for i, (x, y) in enumerate(points):
        for j, z in enumerate((z_top, z_top - 5)):
            row = ["EmbeddedBeam\\_1\\_1", 500 + 2 * i + j, 1, x, y, z]
            row += [-n, -n - 1, 0.0] + [0.0, 0.0, 0.0] * 3 + [100.0, 99.0, 101.0] + [0.0, 0.0, 0.0]
            rows.append(row + [1.0, 2.0, 3.0, "N/A"])
    return rows


def test_slab_design_with_zones_punching_and_restraint():
    def forces(x, y):
        # N_1, N_2, Q_12, Q_23, Q_13, M_11, M_22, M_12: sagging mid-span, hogging over the piles at x = -4.
        near = math.hypot(x + 4, y) < 1.5
        mx = -600.0 if near else 150.0
        return [0.0, 0.0, 0.0, 20.0, 30.0, mx, 80.0, 0.0]

    pts = [(-4.0, 0.0)]
    raw = {
        "Deck-PT-B-Apron": deck_rows(forces),
        "Deck-QP": deck_rows(lambda x, y: [0.0] * 5 + [100.0, 50.0, 0.0]),
        "Pile(1)-PT-B-Apron": piles_at(pts),
        "Pile(1)-QP": piles_at(pts, 1000.0),
    }
    wb = import_sheets(raw)
    els = {
        "Deck": SlabInput(
            thickness=800,
            crack_width_limit=0.3,
            crack_width_limit_bottom=0.3,
            peaks="design",
            restraint_check="design",
        ),
        "Pile(1)": PileInput(head_level=2.7),
    }
    res = run_section(DesignSettings(), Section(elements=els), wb)
    (d,) = res["slabs"]
    assert d["passed"], d["notes"]
    top = d["layers"]["top_x"]
    # Hogging over the pile needs heavier top bars there than the basic mesh.
    assert top["zones"] and all(z["as_mm2_per_m"] > top["basic"]["as_mm2_per_m"] for z in top["zones"])
    assert all(-6 <= z["x"][0] and z["x"][1] <= -2 for z in top["zones"])
    (p,) = d["punching"]
    assert (p["pile"], p["x"], p["y"], p["V_kN"]) == ("Pile(1)", -4.0, 0.0, 1500.0)
    assert p["u1_mm"] == round(math.pi * (1200 + 4 * p["d_mm"]))
    assert p["beta"] == pytest.approx(
        1 + 0.6 * math.pi * 100 / 1500 * 1000 / (1200 + 4 * p["d_mm"]), abs=1e-3
    )
    # At the pile face the β of u1 (EC2 6.4.5(3)); the kmax = 1.5 limit on links.
    assert p["vEd_face_MPa"] == pytest.approx(p["beta"] * 1500e3 / (math.pi * 1200 * p["d_mm"]), rel=2e-3)
    assert p["kmax_ratio"] == pytest.approx(p["vEd_MPa"] / (1.5 * p["vRd_c_MPa"]), abs=2e-3)
    # Restraint from the joints works along the quay: on the bars along Y when the strips run along X.
    assert set(d["restraint"]["layers"]) == {"bottom_y", "top_y"}
    assert d["steel"]["kg_per_m3"] > 0 and d["bands"] and len(d["bands"][0]) == 5


def test_zones_split_where_bars_change_but_not_into_short_pieces():
    options = bar_options(DesignSettings())
    idx = [3, 3, 3, 7, 3, 3] + [0] * 18
    cells = pd.DataFrame({"i": list(range(6)) * 4, "j": [0] * 6 + [1] * 6 + [2] * 6 + [3] * 6, "idx": idx})
    ok = np.zeros((len(idx), len(options)), bool)
    for c, k in enumerate(idx):
        ok[c, k:] = True
    z = zones_for(cells, ok, options, 1.0, 0.0, 0.0, "X", 2.0)
    # The 1 m piece needing option 7 joins its neighbour; the rest keeps option 3.
    assert [(q["x"], q["as_mm2_per_m"]) for q in z["zones"]] == [
        ([0.0, 4.0], round(options[7][0])),
        ([4.0, 6.0], round(options[3][0])),
    ]
    assert list(z["cell_index"][:6]) == [7, 7, 7, 7, 3, 3] and set(z["cell_index"][6:]) == {0}


def deck_workbook():
    def forces(x, y):
        mx = -600.0 if math.hypot(x + 4, y) < 1.5 else 150.0
        return [0.0, 0.0, 0.0, 20.0, 30.0, mx, 80.0, 0.0]

    pts = [(-4.0, 0.0)]
    return import_sheets(
        {
            "Deck-PT-B-Apron": deck_rows(forces),
            "Deck-QP": deck_rows(lambda x, y: [0.0] * 5 + [100.0, 50.0, 0.0]),
            "Pile(1)-PT-B-Apron": piles_at(pts),
            "Pile(1)-QP": piles_at(pts, 1000.0),
        }
    )


def design_deck(**slab):
    limits = {"crack_width_limit": 0.3, "crack_width_limit_bottom": 0.3, "peaks": "design"} | slab
    els = {"Deck": SlabInput(thickness=800, **limits), "Pile(1)": PileInput(head_level=2.7)}
    return run_section(DesignSettings(), Section(elements=els), deck_workbook())["slabs"][0]


def test_slab_user_meshes_punching_depth_crane_and_peaks():
    plain = design_deck()
    d = design_deck(
        mesh_top_x=SlabMesh(diameter=20, spacing=150),
        punching_depths=[PunchingDepth(x=-4.0, y=0.0, thickness=900)],
        crane=[CraneArea(x_from=-8, x_to=-6, y_from=-4, y_to=4, my=900.0)],
    )
    top = d["layers"]["top_x"]
    assert top["basic"]["label"] == "Ø20 @ 150" and top["basic"]["set_by"] == "user"
    assert plain["layers"]["top_x"]["basic"]["set_by"] == "least steel"
    (p,) = d["punching"]
    assert p["thickness_mm"] == 900 and p["thickness_from"] == "set for this pile"
    assert p["d_mm"] == plain["punching"][0]["d_mm"] + 100
    assert p["r_u1_mm"] == 600 + 2 * p["d_mm"]
    # The crane adds sagging My over X -8 to -6: heavier bottom bars along Y there only.
    zones = d["layers"]["bottom_y"]["zones"]
    assert zones and all(z["x"][0] >= -8 and z["x"][1] <= -5 for z in zones)
    assert not plain["layers"]["bottom_y"]["zones"]
    assert any("Mobile crane" in n for n in d["notes"])
    # Averaging the hogging over a ring round the pile needs less top steel there.
    avg = design_deck(peaks="average")
    most = lambda r: max([z["as_mm2_per_m"] for z in r["layers"]["top_x"]["zones"]] or [0])  # noqa: E731
    assert most(avg) < most(plain)


def test_mesh_with_additional_bars_or_mesh_only():
    d = design_deck()
    top = d["layers"]["top_x"]
    assert top["mode"] == "mesh_and_additional" and top["zones"]
    # Zones add bars between the mesh bars; the mesh itself stays everywhere.
    assert all(0 < z["additional_mm2_per_m"] < z["as_mm2_per_m"] for z in top["zones"])
    assert all(
        z["as_mm2_per_m"] - z["additional_mm2_per_m"] == top["basic"]["as_mm2_per_m"] for z in top["zones"]
    )
    only = design_deck(layout_top_x="mesh_only")["layers"]["top_x"]
    assert only["mode"] == "mesh_only" and not only["zones"] and only["utilisation"] <= 1
    assert only["basic"]["as_mm2_per_m"] >= max(z["as_mm2_per_m"] for z in top["zones"]) * 0.8


def test_office_beta_at_the_pile_face():
    ec2 = design_deck()["punching"][0]
    (p,) = design_deck(punching_face_beta="office")["punching"]
    beta0 = 1 + 0.6 * math.pi * 100 / 1500 * 1000 / 1200
    assert p["vEd_face_MPa"] == pytest.approx(beta0 * 1500e3 / (math.pi * 1200 * p["d_mm"]), rel=2e-3)
    assert p["vEd_face_MPa"] > ec2["vEd_face_MPa"]


def test_automatic_stations_round_the_pile_rows():
    # The sample's rows 6.5, 11 and 17 m from the front beam, slab 22.2 m: 8.5 and 9 meet halfway.
    assert auto_stations([6.5, 11.0, 17.0], 22.2) == [0.0, 4.5, 8.75, 13.0, 15.0, 19.0, 22.2]
    assert auto_stations([0.5], 10.0) == [0.0, 2.5, 10.0]


def test_column_and_field_strips_by_station():
    d = design_deck()
    sd = d["strip_design"]
    assert d["strips"] == "column_and_field" and sd["lines"] == [0.0] and sd["column_width_m"] == 2.2
    # No front beam here: stations from the X = 0 edge, 2 m each side of the pile row at 4 m.
    assert sd["stations"] == [0.0, 2.0, 6.0, 8.0]
    rows = {(r["layer"], tuple(r["station"]), r["strip"]): r for r in sd["rows"]}
    col = rows[("top_x", (2.0, 6.0), "column")]
    # Every node of the column strip at the pile row (outside the pile) hogs 600 kNm/m.
    assert col["M_kNm_per_m"] == pytest.approx(600.0) and col["moment"] == "M11"
    assert col["MRd_kNm_per_m"] >= 600 and col["ratio"] <= 1 and col["wk_mm"] <= 0.3
    field = rows[("top_x", (2.0, 6.0), "field")]
    assert field["M_kNm_per_m"] < 600 and field["as_mm2_per_m"] <= col["as_mm2_per_m"]
    assert {(r["moment"], r["strip"]) for r in sd["summary"]} == {("M11", "column"), ("M11", "field")}
    # M22 is not split into strips: one mesh over the whole deck, zones only where it needs more.
    m22 = [r for r in sd["table"] if r["moment"] == "M22"]
    assert m22[0]["label"] == "Whole deck, basic mesh" and m22[0]["strip"] == "all"
    assert all(r["zone"] for r in m22[1:]) and set(m22[0]["bars"]) == {"bottom", "top"}
    assert not d["layers"]["bottom_y"]["strips"] and d["layers"]["bottom_x"]["strips"]
    assert all(r["set_by"] and r["bar_layers"] for r in sd["table"])
    across = sd["across_profile"]
    assert across["moment"] == "M22" and across["axis"] == "Y" and across["lines"] == [0.0]
    # Each strip's QP sets carry the crack terms of the bars the strip gets, as its crack check.
    sag = rows[("bottom_x", (2.0, 6.0), "column")]
    q = sag["sets"]["qp"][0]["crack"]
    assert q["wk_mm"] == pytest.approx(sag["wk_mm"], abs=1e-3) and q["face"] == "bottom"
    assert q["util"] == pytest.approx(q["wk_mm"] / q["limit_mm"], abs=5e-3) and 0 < q["x_mm"] < q["h_mm"]
    # The 3D crack view: wk / limit per cell, the column strip's cells at its crack width.
    worst = max(r["wk_mm"] / r["wk_limit_mm"] for r in sd["rows"] if r["wk_mm"] is not None)
    assert max(b[3] for b in d["crack_bands"]) == pytest.approx(worst, abs=5e-3)
    assert design_deck(stations=[3.0, 5.0])["strip_design"]["stations"] == [0.0, 3.0, 5.0, 8.0]
    uniform = design_deck(strips="uniform")
    assert uniform["strip_design"] is None and uniform["crack_bands"]


def test_slab_meshes_at_150_or_200():
    d = design_deck()
    assert {lay["basic"]["spacing_mm"] for lay in d["layers"].values()} <= {150.0, 200.0}
    s = DesignSettings()
    s.reinforcement.slab_spacings = []
    assert {o[2] for o in bar_options(s)} > {150.0, 200.0}


def test_stations_from_the_front_wall_line_on_the_sea_side():
    box = {"X": [-23.2, -1.0], "Y": [-16.8, 16.8]}
    front = {"type": "front_beam", "box": {"X": [-1.0, 1.0], "Y": [-16.8, 16.8]}}
    piles = [(-7.5, 0.0, 0.6), (-12.0, 0.0, 0.6), (-18.0, 0.0, 0.6)]
    f = strip_frame(SlabInput(), box, piles, [front])
    # Station 0 is the front beam's centre line; the slab runs from 1.0 to 23.2 m, rows at 7.5, 12, 18.
    assert (f["origin"], f["sign"], f["start"], f["end"]) == (0.0, -1.0, 1.0, 23.2)
    assert f["rows"] == [7.5, 12.0, 18.0]
    assert f["bounds"] == [1.0, 5.5, 9.75, 14.0, 16.0, 20.0, 23.2]
    assert strip_frame(SlabInput(), box, piles, [front], [4.0, 8.0])["bounds"] == [1.0, 4.0, 8.0, 23.2]


def test_station_labels_as_in_the_report():
    every = [[2.25, 4.0], [4.0, 8.0], [8.0, 12.0], [12.0, 16.0]]
    assert station_text(every, every) == "All stations"
    assert station_text([[4.0, 8.0]], every) == "Station 4 to 8"
    assert station_text([[2.25, 4.0], [8.0, 12.0], [12.0, 16.0]], every) == "All stations except 4 to 8"


def test_report_table_and_bars_set_by_the_user():
    d = design_deck()
    sd = d["strip_design"]
    table = sd["table"]
    m11 = [r for r in table if r["moment"] == "M11"]
    assert all(r["along_strips"] and len(r["stations"]) == 1 for r in m11)
    assert {r["label"] for r in table if r["moment"] == "M22"} <= {"All stations"} | {
        r["label"] for r in table if r["moment"] == "M22"
    }
    assert sd["profile"]["M11"] and {"column_max", "field_min"} <= set(sd["profile"]["M11"][0])
    assert d["moment_cells"]["cells"] and len(d["moment_cells"]["cells"][0]) == 6
    row = next(r for r in m11 if r["strip"] == "column" and r["stations"] == [[2.0, 6.0]])
    labels = d["layers"]["top_x"]["additional_labels"]
    heavy = labels[-1]
    els = {
        "Deck": SlabInput(thickness=800, crack_width_limit=0.3, crack_width_limit_bottom=0.3, peaks="design"),
        "Pile(1)": PileInput(head_level=2.7),
    }
    sec = Section(elements=els)
    sec.slab_strips["Deck"] = SlabStrips(bars={k: heavy for k in row["keys"]["top"]})
    mine = run_section(DesignSettings(), sec, deck_workbook())["slabs"][0]
    got = next(
        r
        for r in mine["strip_design"]["table"]
        if r["moment"] == "M11" and r["strip"] == "column" and r["stations"] == [[2.0, 6.0]]
    )
    assert got["user_set"] and got["additional"]["top"] == heavy
    assert got["ratio"] < row["ratio"]
    # Stations set on the Design tab win over the slab's own.
    sec.slab_strips["Deck"] = SlabStrips(stations=[3.0, 5.0])
    assert run_section(DesignSettings(), sec, deck_workbook())["slabs"][0]["strip_design"]["stations"] == [
        0.0,
        3.0,
        5.0,
        8.0,
    ]


def test_twisting_moment_is_left_out_unless_asked():
    assert any("without the twisting moment" in n for n in design_deck()["notes"])
    assert any("Wood–Armer" in n for n in design_deck(twisting="wood_armer")["notes"])


def test_bar_layers_set_by_the_user_and_the_mesh_across():
    d = design_deck()
    row = next(
        r
        for r in d["strip_design"]["table"]
        if r["moment"] == "M11" and r["strip"] == "column" and r["stations"] == [[2.0, 6.0]]
    )
    els = {
        "Deck": SlabInput(thickness=800, crack_width_limit=0.3, crack_width_limit_bottom=0.3, peaks="design"),
        "Pile(1)": PileInput(head_level=2.7),
    }
    sec = Section(elements=els)
    spec = "layers: Ø32@150 | Ø25@150"
    bars = {k: spec for k in row["keys"]["top"]} | {"top_y|mesh": "Ø20 @ 150"}
    sec.slab_strips["Deck"] = SlabStrips(bars=bars)
    mine = run_section(DesignSettings(), sec, deck_workbook())["slabs"][0]
    got = next(
        r
        for r in mine["strip_design"]["table"]
        if r["moment"] == "M11" and r["strip"] == "column" and r["stations"] == [[2.0, 6.0]]
    )
    assert got["user_set"] and got["additional"]["top"] == spec and got["set_by"]["top"] == "your bars"
    lay = got["bar_layers"]["top"]
    # Layer 1: the mesh with Ø32 between its bars, at the cover; layer 2 under it, deeper in.
    assert [x["layer"] for x in lay] == [1, 2] and "Ø32 @ 150" in lay[0]["text"]
    assert lay[0]["from_face_mm"] == 50 + 16 and lay[1]["from_face_mm"] > lay[0]["from_face_mm"] + 32
    assert got["bars"]["top"].endswith("Ø32 @ 150 between the mesh bars + Ø25 @ 150 layer 2")
    assert got["ratio"] < row["ratio"]
    assert mine["layers"]["top_y"]["basic"]["label"] == "Ø20 @ 150"
    whole = next(r for r in mine["strip_design"]["table"] if r["label"] == "Whole deck, basic mesh")
    assert whole["set_by"]["top"] == "your mesh" and whole["user_set"]


def test_slab_sets_in_the_adsec_force_set_workbook():
    from io import BytesIO

    from openpyxl import load_workbook

    from triton.design.governing import workbook

    d = design_deck()
    ws = load_workbook(BytesIO(workbook("P", "S", {"slabs": [d]})))["Slabs"]
    rows = [r for r in ws.iter_rows(values_only=True) if r[0]]
    titles = [r[0] for r in rows if r[0].startswith("Deck · ")]
    assert len(titles) == len(d["strip_design"]["table"])
    sets = [r for r in rows if r[0].startswith(("QP ", "ULS "))]
    assert sets and all(r[4] in ("bottom", "top") for r in sets)
    assert {r[0].split()[0] for r in sets} == {"QP", "ULS"}


def test_slab_bars_in_the_drawing_file():
    from triton.design.export import pile_cages

    d = design_deck()
    (js,) = pile_cages("P", {"slabs": [d]})["slabs"]
    assert {f["face"] + f["bars_along"] for f in js["faces"]} == {"bottomX", "bottomY", "topX", "topY"}
    zones = [z for f in js["faces"] for z in f["zones"]]
    assert zones and all(z["layers"] and z["layers"][0]["bars"][0]["kind"] == "mesh" for z in zones)
    assert all(f["mesh"]["layers"][0]["from_face_mm"] > f["cover_mm"] for f in js["faces"])


def test_failing_punching_says_what_would_fix_it():
    wb = import_sheets(
        {
            "Deck-PT-B-Apron": deck_rows(lambda x, y: [0.0] * 5 + [50.0, 50.0, 0.0]),
            "Deck-QP": deck_rows(lambda x, y: [0.0] * 5 + [20.0, 20.0, 0.0]),
            "Pile(1)-PT-B-Apron": piles_at([(-4.0, 0.0)], 9000.0),
            "Pile(1)-QP": piles_at([(-4.0, 0.0)], 6000.0),
        }
    )
    els = {"Deck": SlabInput(thickness=500), "Pile(1)": PileInput(head_level=2.7)}
    (d,) = run_section(DesignSettings(), Section(elements=els), wb)["slabs"]
    (p,) = d["punching"]
    assert not p["passed"]
    fix = p["fix"]
    # A thicker slab always helps, and it needs more without links than with them.
    assert 500 < fix["thickness_mm_with_links"] <= fix["thickness_mm_without_links"]
    if fix.get("rho_l_with_links"):
        assert fix["rho_l_with_links"] > p["rho_l"]


def test_slab_is_designed_with_each_mesh_spacing_and_the_pick_drives_it():
    from triton.project import SlabStrips

    els = {
        "Deck": SlabInput(thickness=800, crack_width_limit=0.3, crack_width_limit_bottom=0.3, peaks="design"),
        "Pile(1)": PileInput(head_level=2.7),
    }
    d = run_section(DesignSettings(), Section(elements=els), deck_workbook())["slabs"][0]
    mc = d["mesh_choice"]
    assert [o["spacing_mm"] for o in mc["options"]] == [150, 200] and not mc["picked"]
    assert all(o["kg_per_m3"] > 0 and o["ratio_pct"] > 0 for o in mc["options"])
    assert {lay["basic"]["spacing_mm"] for lay in d["layers"].values()} == {mc["chosen_mm"]}
    other = 350 - mc["chosen_mm"]
    sec = Section(elements=els, slab_strips={"Deck": SlabStrips(spacing=other)})
    d2 = run_section(DesignSettings(), sec, deck_workbook())["slabs"][0]
    assert d2["mesh_choice"]["chosen_mm"] == other and d2["mesh_choice"]["picked"]
    assert {lay["basic"]["spacing_mm"] for lay in d2["layers"].values()} == {other}
    assert "as you picked" in d2["notes"][0]


def test_added_layers_sit_inside_their_mesh():
    from triton.design.export import _slab
    from triton.design.slabs import bar_layers

    rows = bar_layers((0, 20, 150, 1), [(25, 150), (25, 150)], 50)
    assert [r["from_face_mm"] for r in rows] == sorted(r["from_face_mm"] for r in rows)
    d = {"element": "Deck", "thickness_mm": 800, "layers": {}}
    for face in ("bottom", "top"):
        d["layers"][f"{face}_x"] = {"basic": {}, "mesh_bar_layers": rows, "zones": []}
    faces = {f["face"]: f["mesh"]["layers"] for f in _slab(d)["faces"]}
    up = [r["above_soffit_mm"] for r in faces["bottom"]]
    down = [r["above_soffit_mm"] for r in faces["top"]]
    assert up == sorted(up) and down == sorted(down, reverse=True)  # bottom layers go up, top layers down


def test_slab_station_diagrams_have_the_qp_envelope_too():
    sd = design_deck()["strip_design"]
    assert set(sd["profile_qp"]) == set(sd["profile"]) and sd["across_profile_qp"]["points"]
    # QP moments are smaller than ULS ones in this workbook.
    top = max(abs(q["column_max"]) for q in sd["profile"]["M11"])
    assert max(abs(q["column_max"]) for q in sd["profile_qp"]["M11"]) < top
