from conftest import (
    BEAM_HEADER,
    EMBEDDED_HEADER,
    PLATE_HEADER,
    beam_row,
    plate_row,
    plate_sheet,
)

from triton.elements import CombinationType, ElementType, combination_type, parse_sheet_name
from triton.importer import clean_sheet, normalise_header


def codes(sheet, severity=None):
    return {i.code for i in sheet.issues if severity is None or i.severity.value == severity}


def test_parse_sheet_names():
    p = parse_sheet_name("Pile(1)-PT-B-Apron")
    assert (p.element, p.combination, p.spec.type) == ("Pile(1)", "PT-B-Apron", ElementType.PILE)
    p = parse_sheet_name("Combi Wall-QP")
    assert (p.element, p.combination, p.spec.type) == ("Combi Wall", "QP", ElementType.COMBI_WALL)
    assert parse_sheet_name("Front Beam-PT-B-Yard").spec.type is ElementType.FRONT_BEAM
    assert parse_sheet_name("Deck-PT-C-Apron").spec.type is ElementType.SLAB
    assert parse_sheet_name("Summary") is None
    assert parse_sheet_name("Pile(1)") is None


def test_combination_types():
    assert combination_type("QP") is CombinationType.SLS_QP
    assert combination_type("PT-C-Yard") is CombinationType.ULS
    assert combination_type("Seismic-Apron") is CombinationType.SEISMIC
    assert combination_type("ACC-Berthing") is CombinationType.ACCIDENTAL


def test_normalise_header():
    assert normalise_header("       N_1,min [kN/m]") == ("N_1_min", "kN/m")
    assert normalise_header("  N_min [kN]") == ("N_min", "kN")
    assert normalise_header("    M_11 [kN m/m]") == ("M_11", "kN m/m")
    assert normalise_header("     Structural element") == ("plaxis_label", None)
    assert normalise_header(" Node") == ("Node", None)


def test_clean_plate_sheet_drops_blank_and_duplicate_rows():
    rows = [
        PLATE_HEADER,
        plate_row(1),
        [],
        [None, "  ", None],
        plate_row(1, label="Element 11-1108 (Plate)", local=5),  # same node from another element
        plate_row(2, label="(New\\_Deck\\_Cracked)"),
        plate_row(3, label=""),
    ]
    s = clean_sheet("Deck-QP", rows)
    assert len(s.frame) == 3
    assert s.blank_rows == 2 and s.duplicate_rows == 1
    assert {"blank_rows_removed", "duplicate_rows_removed"} <= codes(s, "info")
    assert not codes(s, "error") and not codes(s, "warning")
    assert list(s.frame["Node"]) == [1, 2, 3]
    assert s.units["M_11"] == "kN m/m"


def test_combi_wall_second_header_is_merged():
    rows = [BEAM_HEADER, beam_row(1, z=2.7), beam_row(2, z=-10.0), [], [], EMBEDDED_HEADER]
    rows += [beam_row(3, "EmbeddedBeam\\_2\\_1", z=-30.0, extras=[1.0, 2.0, 3.0, "   N/A"])]
    s = clean_sheet("Combi Wall-PT-B-Apron", rows)
    assert not codes(s, "error"), s.issues
    assert "repeated_header_removed" in codes(s, "info")
    assert len(s.frame) == 3
    assert "T_skin" in s.frame.columns
    note = next(i for i in s.issues if i.code == "repeated_header_removed")
    assert "F_foot" in note.message and note.rows == [6]


def test_text_in_number_column_is_an_error():
    bad = plate_row(2)
    bad[8] = "12,5 kN"
    s = clean_sheet("SPW-QP", [PLATE_HEADER, plate_row(1), bad])
    assert "non_numeric" in codes(s, "error")
    assert next(i for i in s.issues if i.code == "non_numeric").rows == [3]


def test_missing_required_column():
    header = PLATE_HEADER[:-3]  # no M_12
    rows = [header] + [plate_row(1)[:-3]]
    s = clean_sheet("SPW-QP", rows)
    assert "missing_columns" in codes(s, "error")


def test_empty_sheet_and_missing_header():
    assert clean_sheet("Portal Frame", []).empty
    s = clean_sheet("SPW-QP", [["just", "text"], [1, 2, 3]])
    assert "no_header" in codes(s, "error")


def test_value_outside_envelope_is_flagged():
    row = plate_row(1)
    row[6] = 999.0  # N_1 above its N_1,max
    s = clean_sheet("SPW-QP", [PLATE_HEADER, row])
    assert "outside_envelope" in codes(s, "warning")


def test_node_with_two_coordinates_is_an_error():
    s = clean_sheet("SPW-QP", [PLATE_HEADER, plate_row(1, z=0.0), plate_row(1, z=-1.0)])
    assert "node_coordinates_differ" in codes(s, "error")


def test_embedded_beam_skin_forces_do_not_create_duplicates():
    a = beam_row(1, "EmbeddedBeam\\_1\\_1", extras=[1.0, 2.0, 3.0, 5.0])
    b = beam_row(1, "EmbeddedBeam\\_1\\_1", extras=[1.5, 2.5, 3.5, 5.0])
    s = clean_sheet("Pile(1)-QP", [EMBEDDED_HEADER, a, b])
    assert len(s.frame) == 1
    assert "node_values_differ" not in codes(s)


def test_same_node_different_forces_is_a_warning():
    s = clean_sheet("SPW-QP", [PLATE_HEADER, plate_row(1), plate_row(1, scale=2.0)])
    assert "node_values_differ" in codes(s, "warning")


def test_unexpected_units():
    header = list(PLATE_HEADER)
    header[6] = "N_1 [MN/m]"
    s = clean_sheet("SPW-QP", [header] + plate_sheet()[1:])
    assert "unexpected_units" in codes(s, "warning")


def test_values_in_columns_without_a_header_are_a_quiet_note_naming_the_column():
    rows = [PLATE_HEADER, plate_row(1), plate_row(2) + [None, "max"], plate_row(3)]
    s = clean_sheet("Deck-QP", rows)
    (i,) = [i for i in s.issues if i.code == "unnamed_columns"]
    assert i.severity.value == "info" and i.rows == [3]
    from openpyxl.utils import get_column_letter

    col = get_column_letter(len(PLATE_HEADER) + 2)
    assert (
        i.message == f"Extra column {col} beside the Plaxis table (no header above it): not read. "
        "1 row(s) have values there, e.g. 'max'."
    )
    assert len(s.frame) == 3


def test_xlsx_rows_count_from_excel_row_1_whatever_size_the_file_records(tmp_path):
    from openpyxl import Workbook

    from triton.reader import read_workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Deck-QP"
    for n, r in enumerate([PLATE_HEADER, plate_row(1), plate_row(2)], start=3):
        for c, v in enumerate(r, start=1):
            ws.cell(n, c, v)
    path = tmp_path / "w.xlsx"
    wb.save(path)
    rows = read_workbook(path)["Deck-QP"]
    assert rows[2][: len(PLATE_HEADER)] == PLATE_HEADER  # Excel row 3
    assert not any(rows[0]) and not any(rows[1])
