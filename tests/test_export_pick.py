"""Exports of only the elements picked (?elements=A,B), not the whole section."""

from __future__ import annotations

import io
import zipfile

from conftest import pile_sheet
from openpyxl import load_workbook
from test_design import xlsx_bytes


def workbook_text(content: bytes) -> str:
    wb = load_workbook(io.BytesIO(content))
    return " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)


def test_every_export_takes_the_picked_elements(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from triton.api import app

    monkeypatch.setenv("TRITON_DATA_DIR", str(tmp_path))
    client = TestClient(app)
    names = ["Pile(1)", "Pile(2)"]
    p = client.post("/api/projects", json={"info": {"name": "Berth 1"}, "element_names": names}).json()
    url = f"/api/projects/{p['id']}/sections/{p['sections'][0]['id']}"
    sheets = {f"{n}-{c}": pile_sheet() for n in names for c in ("PT-B-Apron", "QP")}
    client.post(f"{url}/workbook", files={"file": ("s.xlsx", xlsx_bytes(sheets))})
    assert client.post(f"{url}/design").status_code == 200

    def ads(**params):
        r = client.get(f"{url}/design/adsec.zip", params=params)
        assert r.status_code == 200
        return r, zipfile.ZipFile(io.BytesIO(r.content)).namelist()

    _, every = ads()
    r, one = ads(elements="Pile(2)")
    assert any("Pile(1)" in f for f in every) and any("Pile(2)" in f for f in every)
    assert one and all("Pile(2)" in f for f in one)
    assert "Section_1_Pile_2-adsec.zip" in r.headers["content-disposition"]
    assert ads(elements="Pile(1),Pile(2)")[1] == every

    def cells(**params):
        return workbook_text(client.get(f"{url}/design/governing.xlsx", params=params).content)

    text = cells(elements="Pile(1)")
    assert "Pile(1)" in text and "Pile(2)" not in text
    assert "Pile(2)" in cells()

    cages = client.get(f"{url}/design/cages.json", params={"elements": "Pile(1)"}).json()
    assert [c["element"] for c in cages["piles"]] == ["Pile(1)"]
    views = client.get(f"{url}/design/drawings.json", params={"elements": "Pile(2)"}).json()["views"]
    assert views and all(v["element"] == "Pile(2)" for v in views)
    rep = client.get(f"{url}/design/report.xlsx", params={"elements": "Pile(2)"})
    assert rep.status_code == 200 and "Pile_2_summary.xlsx" in rep.headers["content-disposition"]
    text = workbook_text(rep.content)
    assert "Pile(2)" in text and "Pile(1)" not in text

    r = client.get(f"{url}/design/adsec.zip", params={"elements": "Deck"})
    assert r.status_code == 404 and "Deck" in r.json()["detail"]
