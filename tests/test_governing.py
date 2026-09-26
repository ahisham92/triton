import io

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from test_curtailment import LOADS
from test_design import pile_sheets

from triton.design.governing import STEEL_COLUMNS, STEEL_KEYS, pick_sets, steel_sets, workbook
from triton.design.piles import design_pile
from triton.project import DesignSettings, PileInput


def frame():
    return pd.DataFrame(
        {
            "combination": ["A", "A", "B", "B"],
            "Node": [1, 2, 3, 4],
            "Z": [0.0, -1.0, 0.0, -1.0],
            "N": [100.0, 400.0, -50.0, 250.0],
            "M_2": [10.0, -30.0, 5.0, 20.0],
            "M_3": [500.0, 100.0, -700.0, 300.0],
        }
    )


def test_seven_sets_with_corresponding_actions():
    rows = pick_sets(frame(), np.array([0.2, 0.9, 0.5, 0.3]), "most utilised")
    by = {r["case"]: r for r in rows}
    cases = ["max N", "min N", "max M2", "min M2", "max M3", "min M3", "most utilised"]
    assert [r["case"] for r in rows] == cases
    assert (by["max N"]["node"], by["max N"]["M2_kNm"], by["max N"]["M3_kNm"]) == (2, -30.0, 100.0)
    assert by["min N"]["node"] == 3
    assert by["max M2"]["node"] == 4 and by["min M2"]["node"] == 2
    assert by["max M3"]["node"] == 1 and by["min M3"]["node"] == 3
    # The most utilised point is also max N: the row is kept twice.
    assert by["most utilised"]["node"] == 2 and by["most utilised"]["utilisation"] == 0.9
    # Without utilisations (QP) the 7th set is the largest resultant moment.
    assert pick_sets(frame(), None, "largest resultant M")[-1]["node"] == 3


def test_sets_per_station_of_a_curtailed_pile():
    qp = [(n, z, 0.5 * nn, 0.5 * m2, m3) for n, z, nn, m2, m3 in LOADS]
    d = design_pile("Pile(1)", PileInput(head_level=0.0), DesignSettings(), pile_sheets(LOADS, qp)).to_dict()
    runs = d["curtailment"]["runs"]
    sets = d["governing_sets"]
    assert [(s["top"], s["bottom"]) for s in sets] == [(r["top"], r["bottom"]) for r in runs]
    top = sets[0]
    assert len(top["uls"]) == 7 and len(top["qp"]) == 7
    assert all(top["bottom"] - 1e-9 <= r["z"] <= top["top"] + 1e-9 for r in top["uls"] + top["qp"])
    # Concrete sign: Plaxis N = -3000 is +3000 here; QP is half.
    assert top["uls"][0]["N_kN"] == pytest.approx(3000.0)
    assert top["qp"][0]["N_kN"] == pytest.approx(1500.0)
    # The head station's most utilised point is the pile's governing point.
    assert top["uls"][-1]["utilisation"] == pytest.approx(d["utilisation"], abs=1e-3)
    # QP sets: SLS, the crack width over its limit.
    assert top["qp"][-1]["utilisation"] == top["qp"][-1]["crack"]["util"]
    # Each QP set carries its crack width and the terms the crack picture draws.
    for r in (r for st in sets for r in st["qp"]):
        c = r["crack"]
        assert c["h_mm"] == 1200 and c["limit_mm"] == 0.2
        assert c["util"] == pytest.approx(c["wk_mm"] / 0.2, abs=5e-3)
        assert (c["wk_mm"] == 0) == (c["sr_max_mm"] is None)
        assert c["wk_mm"] > 0 or c["x_mm"] == 1200  # no crack: the whole section in compression
    worst = max(r["crack"]["wk_mm"] for st in sets for r in st["qp"])
    assert worst <= d["cracks"]["wk_mm"] + 1e-3
    assert max(b[3] for b in d["cracks"]["bands"]) == pytest.approx(d["cracks"]["wk_mm"] / 0.2, abs=5e-3)


def test_governing_sets_workbook():
    d = design_pile("Pile(1)", PileInput(head_level=0.0), DesignSettings(), pile_sheets(LOADS)).to_dict()
    tube = {"element": "Combi Wall", "infill": {"governing_sets": []}, "tube": {"governing_sets": steel()}}
    spw = {"element": "SPW", "governing_sets": steel("plate")}
    data = workbook("Berth", "Section 01a", {"piles": [d], "combi_walls": [tube], "sheet_pile_walls": [spw]})
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Concrete", "Steel"]
    rows = list(wb["Concrete"].iter_rows(min_row=3, values_only=True))
    stations = d["governing_sets"]
    # Per station: a blank row, the name, the header, 7 QP rows then 7 ULS rows.
    assert len(rows) == 17 * len(stations)
    assert rows[1][0].startswith("Pile(1) · 0 to ")
    assert rows[2][:5] == ("Criterion", "N (kN)", "M2 (kNm)", "M3 (kNm)", "Combination")
    assert rows[3][0] == "QP max N" and rows[10][0] == "ULS max N" and rows[16][0] == "ULS most utilised"
    assert rows[10][1:5] == tuple(
        stations[0]["uls"][0][k] for k in ("N_kN", "M2_kNm", "M3_kNm", "combination")
    )
    steel_rows = list(wb["Steel"].iter_rows(min_row=3, values_only=True))
    assert [r[0] for r in steel_rows if r and r[0] and not r[0].startswith(("ULS", "Criterion"))] == [
        "Combi Wall tube (kN, kNm)",
        "SPW (kN/m, kNm/m)",
    ]
    assert steel_rows[2][:6] == ("Criterion", "N", "M2", "M3", "Q1 (Q_12)", "Q2 (Q_13)")
    assert sum(1 for r in steel_rows if r and r[0] and r[0].startswith("ULS")) == 20


def steel(kind="beam"):
    cols = STEEL_COLUMNS[kind]
    f = pd.DataFrame(
        {
            "combination": ["A", "A", "B", "QP"],
            "Node": [1, 2, 3, 4],
            "Z": [0.0, -1.0, -2.0, -3.0],
            **{c: [float(i + 1), -float(i + 1), 2.0 * (i + 1), 0.5] for i, c in enumerate(cols)},
        }
    )
    return steel_sets(f[f["combination"] != "QP"], kind)


def test_ten_steel_sets_over_all_combinations():
    sets = steel()
    assert [r["case"] for r in sets["rows"]] == [f"{m} {k}" for k in STEEL_KEYS for m in ("max", "min")]
    by = {r["case"]: r for r in sets["rows"]}
    # Every max comes from combination B (node 3), every min from A (node 2), not one per combination.
    assert {r["node"] for c, r in by.items() if c.startswith("max")} == {3}
    assert by["min Q2"]["node"] == 2 and by["min Q2"]["Q2"] == -5.0 and by["min Q2"]["N"] == -1.0
    assert steel("plate")["columns"] == dict(
        zip(STEEL_KEYS, ("N_1", "M_11", "M_22", "Q_13", "Q_23"), strict=True)
    )


def test_single_station_when_not_curtailed():
    settings = DesignSettings()
    settings.piles.curtail = False
    sheets = pile_sheets([(1, 0.0, -2000.0, 1500.0, 0.0), (2, -5.0, -2500.0, 800.0, 0.0)])
    d = design_pile("Pile(1)", PileInput(head_level=0.0), settings, sheets).to_dict()
    (st,) = d["governing_sets"]
    assert (st["top"], st["bottom"], st["cage"]) == (0.0, -5.0, d["arrangement"]["label"])


def test_no_crack_check_gives_qp_rows_of_ones():
    from triton.project import Casing

    pile = PileInput(head_level=0.0, casing=Casing(top_level=0.0, bottom_level=-30.0))
    d = design_pile("Pile(1)", pile, DesignSettings(), pile_sheets(LOADS)).to_dict()
    for st in d["governing_sets"]:
        assert len(st["qp"]) == 7
        assert {(r["N_kN"], r["M2_kNm"], r["M3_kNm"]) for r in st["qp"]} == {(1.0, 1.0, 1.0)}
        assert st["uls"][0]["N_kN"] == pytest.approx(3000.0)
    # A casing over the top 4 m only: stations below it keep their QP sets.
    pile = PileInput(head_level=0.0, casing=Casing(top_level=0.0, bottom_level=-4.0))
    d = design_pile("Pile(1)", pile, DesignSettings(), pile_sheets(LOADS)).to_dict()
    assert d["governing_sets"][-1]["qp"][0]["N_kN"] == pytest.approx(0.6 * 3000.0)
    # Points inside the casing are left out of every station's QP sets.
    qp = [r for st in d["governing_sets"] for r in st["qp"] if r["z"] is not None]
    assert qp and all(r["z"] < -4.0 for r in qp)


def test_spw_export(tmp_path, monkeypatch):
    from conftest import plate_sheet
    from fastapi.testclient import TestClient
    from test_design import xlsx_bytes

    from triton.api import app

    monkeypatch.setenv("TRITON_DATA_DIR", str(tmp_path))
    client = TestClient(app)
    p = client.post("/api/projects", json={"info": {"name": "Berth"}, "element_names": ["SPW"]}).json()
    sid = p["sections"][0]["id"]
    url = f"/api/projects/{p['id']}/sections/{sid}"
    assert client.get(f"{url}/spw.xlsx").status_code == 409
    # A 1.35 multiplier on the Set B sheet; SPW keeps the Plaxis sign of N.
    p["sections"][0]["load_factors"] = [{"factor": 1.35, "sheets": ["SPW-PT-B-Apron"]}]
    assert client.put(f"/api/projects/{p['id']}", json=p).status_code == 200
    data = xlsx_bytes({"SPW-PT-B-Apron": plate_sheet(scale=2.0), "SPW-QP": plate_sheet()})
    client.post(f"{url}/workbook", files={"file": ("s.xlsx", data)})
    r = client.get(f"{url}/spw.xlsx")
    assert (
        r.status_code == 200
        and "Berth_Section_1_SPW-straining-actions.xlsx" in r.headers["content-disposition"]
    )
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Durability", "Governing", "Envelope PT-B-Apron", "Envelope QP"]
    cells = [r for r in wb["Durability"].iter_rows(values_only=True) if r and r[0] is not None]
    assert ("Zone", "Z bottom (m)", "Loss front (mm)", "Loss back (mm)") == cells[3][:4]
    assert cells[4][:4] == (1, -0.5, 4.5, 0)
    rows = list(wb["Governing"].iter_rows(min_row=4, values_only=True))
    header, rows = rows[0], rows[1:]
    assert header[:2] == ("Criterion", "N (N_1)") and header[6] == "Combination"
    # Ten ULS rows over the whole wall; the QP sheet is not a ULS combination.
    assert len(rows) == 10 and {r[6] for r in rows} == {"PT-B-Apron"}
    env = list(wb["Envelope QP"].iter_rows(min_row=4, values_only=True))
    env_b = list(wb["Envelope PT-B-Apron"].iter_rows(min_row=4, values_only=True))
    assert env_b[0][1] == pytest.approx(env[0][1] * 2.0 * 1.35, rel=1e-3)
    assert rows[0][1] == pytest.approx(max(r[1] for r in env_b), rel=1e-3)


def test_durability_tables_per_corrosion_zone():
    from types import SimpleNamespace

    import pandas as pd

    from triton.design.spw import durability
    from triton.project import SheetPileInput, SheetPileZone

    # Largest |M| at z = -1 (zone 2) with its own V; the largest |V| sits elsewhere in that zone.
    f = pd.DataFrame(
        {
            "Z": [2.0, 0.0, -1.0, -3.0, -6.0],
            "M_11": [5.0, -20.0, -80.0, 10.0, 30.0],
            "Q_13": [1.0, 60.0, 15.0, -90.0, 40.0],
            "Q_23": [0.0] * 5,
        }
    )
    sheets = {"SPW-PT-B-Apron": SimpleNamespace(frame=f), "SPW-QP": SimpleNamespace(frame=f * 10)}
    wall = SheetPileInput(
        corrosion_zones=[
            SheetPileZone(bottom_level=-0.5, front=4.5),
            SheetPileZone(bottom_level=-5.0, front=2.5, back=1.75),
            SheetPileZone(bottom_level=-8.0, front=1.75, back=1.75),
        ]
    )
    t = durability(sheets, wall)
    assert t["top"] == 2.0
    assert [(r["z"], r["M"], r["V"]) for r in t["max_M"]] == [
        (-0.5, 20.0, 60.0),
        (-5.0, 80.0, 15.0),
        (-8.0, 30.0, 40.0),
    ]
    assert [(r["M"], r["V"]) for r in t["max_V"]][1] == (10.0, 90.0)
    # A peak inside a king pile (the connection) is left out.
    g = pd.concat([f, pd.DataFrame({"Z": [-2.0], "M_11": [0.0], "Q_13": [900.0], "Q_23": [0.0]})])
    g = g.assign(X=0.0, Y=[1.0] * 5 + [3.2])
    peak = {"SPW-PT-B-Apron": SimpleNamespace(frame=g)}
    assert durability(peak, wall)["max_V"][1]["V"] == 900.0
    assert durability(peak, wall, [(0.0, 3.2, 0.813)])["max_V"][1]["V"] == 90.0
    with pytest.raises(ValueError):
        SheetPileInput(corrosion_zones=[SheetPileZone(bottom_level=-5.0), SheetPileZone(bottom_level=-1.0)])
