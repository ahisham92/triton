"""Displacements typed in per section (as received by email), checked against their limits."""

from __future__ import annotations

import io

from conftest import pile_sheet
from openpyxl import load_workbook
from test_design import xlsx_bytes

from triton.project import Displacement


def test_pass_fail():
    assert Displacement(value=-45, limit=50).passed is True
    assert Displacement(value=62, limit=50).passed is False
    assert Displacement(value=62).passed is None


def test_typed_while_locked_without_making_results_out_of_date(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from triton.api import app

    monkeypatch.setenv("TRITON_DATA_DIR", str(tmp_path))
    client = TestClient(app)
    p = client.post("/api/projects", json={"info": {"name": "B"}, "element_names": ["Pile(1)"]}).json()
    url = f"/api/projects/{p['id']}/sections/{p['sections'][0]['id']}"
    data = xlsx_bytes({"Pile(1)-PT-B-Apron": pile_sheet(), "Pile(1)-QP": pile_sheet()})
    client.post(f"{url}/workbook", files={"file": ("s.xlsx", data)})
    client.post(f"{url}/design")
    page = client.get(f"/api/projects/{p['id']}").json()
    page["locked"] = True
    client.put(f"/api/projects/{p['id']}", json=page)
    page = client.get(f"/api/projects/{p['id']}").json()
    page["sections"][0]["displacements"] = [
        {"what": "Front beam, horizontal", "value": 62, "limit": 50, "combination": "SLS", "source": "Email"},
        {"what": "Deck, vertical", "value": -12, "limit": 25},
    ]
    r = client.put(f"/api/projects/{p['id']}", json=page)
    assert r.status_code == 200, r.text
    assert not client.get(f"{url}/design").json().get("changed")
    wb = load_workbook(io.BytesIO(client.get(f"{url}/design/report.xlsx").content))
    text = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Front beam, horizontal" in text and "NOT OK" in text and "3.5 Displacements" in text
