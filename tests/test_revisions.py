"""Issued revisions (a copy of the project as issued, what changed since) and each element's checking."""

from __future__ import annotations

import io
import zipfile

from conftest import pile_sheet
from openpyxl import load_workbook
from test_design import xlsx_bytes

from triton.revisions import next_revision


def test_next_revision():
    assert [next_revision(r) for r in ("P01", "C09", "A", "2", "x-y")] == ["P02", "C10", "B", "3", "x-y.1"]


def _text(content: bytes) -> str:
    wb = load_workbook(io.BytesIO(content))
    return " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)


def test_issue_compare_and_check(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from triton.api import app

    monkeypatch.setenv("TRITON_DATA_DIR", str(tmp_path))
    client = TestClient(app)
    info = {"name": "Berth 1", "checker": "M. Checker", "document_number": "CALC-001"}
    p = client.post("/api/projects", json={"info": info, "element_names": ["Pile(1)"]}).json()
    sid = p["sections"][0]["id"]
    url = f"/api/projects/{p['id']}/sections/{sid}"
    data = xlsx_bytes({"Pile(1)-PT-B-Apron": pile_sheet(), "Pile(1)-QP": pile_sheet()})
    client.post(f"{url}/workbook", files={"file": ("s.xlsx", data)})
    client.post(f"{url}/design")

    # Checking: open while locked, kept when a page saves the project.
    r = client.put(f"{url}/checks/Pile(1)", json={"status": "checked", "by": "M. Checker", "comment": "OK"})
    assert r.status_code == 200 and r.json()["sections"][0]["checks"]["Pile(1)"]["status"] == "checked"
    assert client.put(f"{url}/checks/Deck", json={"status": "checked"}).status_code == 404

    # Issue P01: a copy is kept and the revision in work moves on.
    issued = client.post(f"/api/projects/{p['id']}/revisions", json={"description": "For approval"}).json()
    assert issued["info"]["revision"] == "P02" and issued["revisions"][0]["rev"] == "P01"
    assert issued["revisions"][0]["checked"] == "M. Checker"
    copy = client.get(f"/api/projects/{p['id']}/revisions/P01/project.trt")
    assert copy.status_code == 200 and "triton.json" in zipfile.ZipFile(io.BytesIO(copy.content)).namelist()

    # A page holding the project without the revision does not undo it.
    page = client.get(f"/api/projects/{p['id']}").json()
    page["revisions"] = []
    page["sections"][0]["checks"] = {}
    page["locked"] = False
    saved = client.put(f"/api/projects/{p['id']}", json=page).json()
    assert (
        saved["revisions"][0]["rev"] == "P01"
        and saved["sections"][0]["checks"]["Pile(1)"]["by"] == "M. Checker"
    )

    # Change the pile and design again: the change and the new results show against P01.
    saved["sections"][0]["elements"]["Pile(1)"]["diameter"] = 1500
    client.put(f"/api/projects/{p['id']}", json=saved)
    client.post(f"{url}/design")
    ch = client.get(f"/api/projects/{p['id']}/revisions/P01/changes").json()
    assert any("diameter" in r["what"] and r["now"] == 1500 for r in ch["inputs"])
    assert ch["sections"][0]["elements"][0]["element"] == "Pile(1)"

    # The report prints the title block, the revisions and the checking (on an earlier design now).
    rep = _text(client.get(f"{url}/design/report.xlsx").content)
    assert "CALC-001" in rep and "P02" in rep and "For approval" in rep and "on an earlier design" in rep

    again = client.post(f"/api/projects/{p['id']}/revisions", json={}).json()
    assert again["info"]["revision"] == "P03"
    page = client.get(f"/api/projects/{p['id']}").json()
    page["info"]["revision"] = "P01"
    client.put(f"/api/projects/{p['id']}", json=page)
    assert client.post(f"/api/projects/{p['id']}/revisions", json={}).status_code == 409
