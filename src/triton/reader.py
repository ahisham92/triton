"""Read raw cell values from .xlsb / .xlsx / .xlsm workbooks.

The reader does no interpretation: each sheet becomes a list of rows, each row
a list of cell values (``None`` for empty cells). Row index 0 is Excel row 1.
"""

from __future__ import annotations

import re
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

Row = list[Any]
# Called after each sheet with the share of the workbook read so far (0..1) and the sheet's name.
Progress = Callable[[float, str], None]


class UnsupportedWorkbook(ValueError):
    pass


def read_workbook(path: str | Path, progress: Progress | None = None) -> dict[str, list[Row]]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xlsb":
        return _read_xlsb(path, progress)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_openpyxl(path, progress)
    raise UnsupportedWorkbook(f"Unsupported file type '{path.suffix}'. Use .xlsb, .xlsx or .xlsm.")


def _sheet_weights(path: Path, count: int) -> list[float]:
    """Each sheet's share of the reading: the size of its part in the file (sheet1, sheet2, … in
    workbook order), or equal shares when the parts cannot be matched to the sheets."""
    try:
        with zipfile.ZipFile(path) as z:
            parts = [
                (int(m.group(1)), i.file_size)
                for i in z.infolist()
                if (m := re.fullmatch(r"xl/worksheets/sheet(\d+)\.(?:bin|xml)", i.filename))
            ]
    except (OSError, zipfile.BadZipFile):
        parts = []
    sizes = [size for _, size in sorted(parts)]
    if len(sizes) != count or not sum(sizes):
        sizes = [1] * count
    total = sum(sizes) or 1
    return [size / total for size in sizes]


class _Tracker:
    def __init__(self, path: Path, names: list[str], progress: Progress | None) -> None:
        self.weights = dict(zip(names, _sheet_weights(path, len(names)), strict=True)) if progress else {}
        self.progress, self.done = progress, 0.0

    def __call__(self, name: str) -> None:
        if self.progress:
            self.done += self.weights.get(name, 0.0)
            self.progress(min(self.done, 1.0), name)


def _read_xlsb(path: Path, progress: Progress | None = None) -> dict[str, list[Row]]:
    from pyxlsb import open_workbook

    sheets: dict[str, list[Row]] = {}
    with open_workbook(str(path)) as wb:
        tick = _Tracker(path, list(wb.sheets), progress)
        for name in wb.sheets:
            rows: list[Row] = []
            with wb.get_sheet(name) as sheet:
                for cells in sheet.rows(sparse=True):
                    if not cells:
                        continue
                    r = cells[0].r
                    while len(rows) < r:
                        rows.append([])
                    values: Row = []
                    for cell in cells:
                        while len(values) < cell.c:
                            values.append(None)
                        values.append(cell.v)
                    rows.append(values)
            sheets[name] = rows
            tick(name)
    return sheets


def _read_openpyxl(path: Path, progress: Progress | None = None) -> dict[str, list[Row]]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    try:
        tick = _Tracker(path, [ws.title for ws in wb.worksheets], progress)
        sheets: dict[str, list[Row]] = {}
        for ws in wb.worksheets:
            sheets[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
            tick(ws.title)
        return sheets
    finally:
        wb.close()
