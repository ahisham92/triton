"""Calculation reports of a designed section: a summary or the detailed calculation, as Excel, Word or PDF.

The content is built once as a list of blocks (headings, paragraphs, key-value lists and tables)
from the stored design results, then written out by one renderer per format.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from . import clock
from .figures import slab_stations
from .materials import STEEL_DENSITY
from .project import DesignSettings, Project, Section


@dataclass
class Block:
    kind: str  # "h1", "h2", "h3", "p", "note", "caption", "bullets", "kv", "table", "image"
    text: str = ""
    rows: list[list[Any]] = field(default_factory=list)
    headers: list[str] = field(default_factory=list)
    image: bytes = b""  # PNG, for kind "image"


@dataclass
class Report:
    title: str
    subtitle: str
    blocks: list[Block] = field(default_factory=list)

    def h(self, level: int, text: str) -> None:
        self.blocks.append(Block(f"h{level}", text))

    def p(self, text: str) -> None:
        if text:
            self.blocks.append(Block("p", text))

    def note(self, text: str) -> None:
        if text:
            self.blocks.append(Block("note", text))

    def image(self, png: bytes, caption: str) -> None:
        self.blocks.append(Block("image", image=png))
        self.caption(caption)

    def caption(self, text: str) -> None:
        self.blocks.append(Block("caption", text))

    def bullets(self, items: list[str]) -> None:
        if items:
            self.blocks.append(Block("bullets", rows=[[x] for x in items]))

    def kv(self, pairs: list[tuple[str, Any]]) -> None:
        rows = [[k, _fmt(v)] for k, v in pairs if v is not None and v != ""]
        if rows:
            self.blocks.append(Block("kv", rows=rows))

    def table(self, headers: list[str], rows: list[list[Any]]) -> None:
        if rows:
            self.blocks.append(Block("table", headers=headers, rows=[[_fmt(v) for v in r] for r in rows]))


def _fmt(v: Any, digits: int = 3) -> str:
    if v is None:
        return "–"
    if isinstance(v, bool):
        return "passes" if v else "fails"
    if isinstance(v, float):
        if abs(v) >= 100:
            return f"{v:,.0f}"
        if abs(v) >= 10:
            return f"{v:,.1f}"
        return f"{v:.{digits}g}"
    return str(v)


def _ok(passed: Any) -> str:
    return "passes" if passed else "fails"


KIND = {
    "front_beam": "Front beam",
    "rear_beam": "Rear beam",
    "transverse_beam": "Transverse beam",
    "slab": "Slab",
}


# --- Content ----------------------------------------------------------------------------------------


def build_report(project: Project, section: Section, results: dict, detail: str = "summary") -> Report:
    """The report on the layout of the office design reports (BS 6349 / Eurocode quay walls).

    1 Introduction, 2 Design criteria, 3 Design of sections for strength and serviceability
    (summary tables of moment, capacity, crack width and governing combination per element part,
    then shear, punching and steel quantities). The detailed report adds an appendix with the
    calculation of each element.
    """
    info = project.info
    run_at = clock.show(results.get("run_at", ""))
    r = Report(
        f"{info.name}: {section.name}",
        f"{'Detailed structural calculations' if detail == 'detailed' else 'Structural design summary'}, "
        f"designed {run_at}",
    )
    r.kv(
        [
            ("Project", info.name),
            ("Project number", info.number),
            ("Client", info.client),
            ("Location", info.location),
            ("Section", section.name),
            ("Designed by", info.designer),
            ("Checked by", info.checker),
            ("Report printed", clock.now().strftime("%Y-%m-%d %H:%M")),
        ]
    )
    if results.get("changed"):
        r.p(
            "OUT OF DATE: these results were designed before the following inputs changed: "
            + ", ".join(results["changed"])
            + ". Design the section again for results that match its inputs."
        )
    _introduction(r, project, section, results, detail)
    _criteria(r, project.design, section, results)
    _sections(r, section, results)
    if detail == "detailed":
        r.h(1, "Appendix A. Calculations of each element")
        for p in results.get("piles", []):
            _pile(r, p)
        for w in results.get("combi_walls", []):
            _combi(r, w)
        for b in results.get("beams", []):
            _beam(r, b)
        for s in results.get("slabs", []):
            _slab(r, s)
        for w in results.get("sheet_pile_walls", []):
            _spw(r, w)
    return r


def _elements(res: dict) -> list[str]:
    out = [f"{p['element']} (bored pile)" for p in res.get("piles", [])]
    out += [
        f"{w['element']} (combi wall: steel tube king piles with concrete infill)"
        for w in res.get("combi_walls", [])
    ]
    out += [f"{b['element']} ({KIND.get(b.get('kind'), 'beam').lower()})" for b in res.get("beams", [])]
    out += [f"{s['element']} (deck slab)" for s in res.get("slabs", [])]
    out += [
        f"{w['element']} (steel sheet pile wall"
        + (
            f", {w['design']['section']})"
            if (w.get("design") or {}).get("section")
            else ", straining actions only)"
        )
        for w in res.get("sheet_pile_walls", [])
    ]
    return out


def _introduction(r: Report, project: Project, section: Section, res: dict, detail: str) -> None:
    r.h(1, "1. Introduction")
    r.h(2, "1.1 Purpose of the report")
    r.p(
        f"This report presents the structural {'calculations' if detail == 'detailed' else 'design results'} "
        f"of {section.name} of {project.info.name}. It documents the design assumptions and demonstrates that "
        "the reinforced concrete and steel elements of the section satisfy the Ultimate Limit State (ULS) and "
        "Serviceability Limit State (SLS) requirements, and gives the reinforcement of each element."
    )
    r.h(2, "1.2 Scope of work")
    r.p("The section comprises the following structural elements:")
    r.bullets(_elements(res))
    r.h(2, "1.3 Methodology overview")
    r.bullets(
        [
            "Straining actions: taken from the Plaxis 3D model of the section (node results per load "
            "combination), within the section's working zone, with the load multipliers of the section applied.",
            "Member design: limit state design to BS EN 1992-1-1 and BS EN 1993 with the partial factor approach "
            "of BS 6349. Concrete sections are designed for axial force with biaxial bending (N–M interaction "
            "as in Oasys AdSec), shear and crack width.",
            "Durability: concrete covers, crack width limits and steel corrosion allowances for the design life "
            "in a marine environment, to BS 6349-1-4.",
        ]
    )
    r.h(2, "1.4 Units and datum")
    r.p(
        "All units are SI units, unless stated otherwise. Levels (z) are those of the Plaxis model, in metres."
    )


def _criteria(r: Report, s: DesignSettings, section: Section, res: dict) -> None:
    pf, m, cr, d = s.partial_factors, s.materials, s.cracking, s.durability
    r.h(1, "2. Design criteria")
    r.h(2, "2.1 Design life and design standards")
    r.p(f"The structure is designed for a {s.design_life_years}-year design life, to:")
    r.bullets(
        [
            "BS 6349-1-1, BS 6349-1-2, BS 6349-1-4 and BS 6349-2: Maritime works.",
            "BS EN 1990: Basis of structural design.",
            "BS EN 1992-1-1: Design of concrete structures.",
            "BS EN 1993-1-1, BS EN 1993-1-6 and BS EN 1993-5: Design of steel structures, shells and piling.",
        ]
    )
    r.h(2, "2.2 Material specifications")
    r.kv(
        [
            ("Concrete (piles, slabs, beams)", m.concrete),
            ("Combi wall infill concrete", m.infill_concrete),
            ("Reinforcement", f"{s.reinforcement.grade} to BS 4449"),
            ("Tube and casing steel", m.structural_steel),
            ("Sheet pile steel", m.sheet_pile_steel),
            ("Bars used", ", ".join(f"Ø{x}" for x in s.reinforcement.bar_diameters)),
        ]
    )
    r.h(2, "2.3 Durability")
    r.h(3, "Crack width limits")
    limits = sorted({_limit(x) for x in section.elements.values() if _limit(x) is not None})
    r.p(
        "Crack widths are checked under the quasi-permanent (QP) combinations (BS EN 1990 6.5.3), limit "
        + (" / ".join(f"{v:g}" for v in limits) or "0.2")
        + " mm. No crack width check applies inside a steel casing or tube."
    )
    r.h(3, "Concrete cover")
    cv = d.covers
    r.kv(
        [
            (
                "Covers from",
                "BS 6349-1-4 (project values)" if d.cover_code == "bs6349" else "BS EN 1992-1-1 Table 4.4N",
            ),
            ("Piles", f"{cv.piles:g} mm"),
            ("Combi wall infill", f"{cv.combi_infill:g} mm"),
            ("Slab top / bottom", f"{cv.slab_top:g} / {cv.slab_bottom:g} mm"),
            ("Beams", f"{cv.beams:g} mm"),
        ]
    )
    r.h(3, "Corrosion allowance for structural steel")
    co = d.corrosion
    r.kv(
        [
            (
                "Allowances from",
                "BS 6349-1-4:2021 mean values" if d.corrosion_code == "bs6349" else "BS EN 1993-5 Table 4.2",
            ),
            ("Pile casing", f"{co.casing:g} mm over {s.design_life_years} years"),
            (
                "Combi wall tube",
                "by zone, in the table below"
                if any(
                    getattr(e, "corrosion_zones", None)
                    for e in section.elements.values()
                    if e.kind == "combi_wall"
                )
                else f"{co.combi_tube:g} mm over {s.design_life_years} years",
            ),
            ("Sheet piles, per face", f"{co.sheet_pile_per_face:g} mm over {s.design_life_years} years"),
        ]
    )
    n_tab = 0
    for name, el in section.elements.items():
        zones = getattr(el, "corrosion_zones", None)
        if not zones:
            continue
        top = "top"
        rows = []
        if el.kind == "combi_wall":
            for z in zones:
                rows.append(
                    [
                        f"{top} to {z.bottom_level:g}",
                        z.outside,
                        z.inside,
                        el.tube_thickness - z.outside - z.inside,
                    ]
                )
                top = f"{z.bottom_level:g}"
            n_tab += 1
            r.caption(
                f"Table 2-{n_tab}: Loss of thickness of the {name} tube, over {s.design_life_years} years"
            )
            r.table(["Level (m)", "Outside (mm)", "Inside (mm)", "Remaining wall (mm)"], rows)
            r.p("The last zone continues to the toe.")
        elif el.kind == "sheet_pile_wall":
            for z in zones:
                rows.append([f"{top} to {z.bottom_level:g}", z.front, z.back])
                top = f"{z.bottom_level:g}"
            n_tab += 1
            r.caption(f"Table 2-{n_tab}: Loss of thickness of the {name}, over {s.design_life_years} years")
            r.table(["Level (m)", "Front face (mm)", "Back face (mm)"], rows)
    r.h(2, "2.4 Partial factors and design assumptions")
    r.kv(
        [
            ("γc / γs", f"{pf.gamma_c:g} / {pf.gamma_s:g}"),
            ("γc / γs accidental and seismic", f"{pf.gamma_c_accidental:g} / {pf.gamma_s_accidental:g}"),
            ("γM0 / γM1", f"{pf.gamma_m0:g} / {pf.gamma_m1:g}"),
            ("αcc", f"{pf.alpha_cc:g}"),
            ("Concrete area", "net of the bars" if pf.deduct_bar_area else "gross (as AdSec)"),
            (
                "Plate moments",
                "positive read from the results (Auto; see each slab and beam)"
                if s.plate_positive_moment == "auto"
                else f"positive = {s.plate_positive_moment}",
            ),
            ("Shear checked at", f"{s.shear_check_distance} from the support face"),
            (
                "Beam bending and shear",
                "peak × beam width" if s.beam_actions == "peak_width" else "integrated over the model width",
            ),
            ("Creep coefficient φ (QP stresses)", f"{cr.creep_coefficient:g}"),
            (
                "Restraint: T1, T2, α, K1",
                f"{cr.early_age_drop:g} °C, {cr.seasonal_drop:g} °C, {cr.thermal_expansion:g} µε/°C, "
                f"{cr.creep_factor:g}",
            ),
        ]
    )
    r.h(2, "2.5 Limit states and load combinations")
    r.p(
        "The load combinations are those of the Plaxis workbook, one sheet per element and combination. ULS "
        "combinations design the sections; QP combinations are used for crack widths only. N is taken "
        "compression positive (Plaxis N × −1) for concrete elements and with the Plaxis sign for steel."
    )
    combos = sorted(_combinations(res))
    if combos:
        r.p("Combinations governing at least one element: " + ", ".join(combos) + ".")
    factors = getattr(section, "load_factors", None) or []
    if factors:
        r.table(
            ["Load multiplier", "Sheets"],
            [[f"× {f.factor:g}", ", ".join(f.sheets)] for f in factors if getattr(f, "sheets", None)],
        )


def _limit(element: Any) -> float | None:
    return getattr(element, "crack_width_limit", None)


def _combinations(res: dict) -> set[str]:
    out: set[str] = set()

    def walk(x: Any) -> None:
        if isinstance(x, dict):
            c = x.get("combination")
            if isinstance(c, str) and c and c != "no crack check":
                out.add(c)
            for v in x.values():
                walk(v)
        elif isinstance(x, list):
            for v in x:
                walk(v)

    walk({k: v for k, v in res.items() if k != "skipped"})
    return out


def _sections(r: Report, section: Section, res: dict) -> None:
    r.h(1, "3. Design of sections for strength and serviceability")
    r.h(2, "3.1 General")
    r.p(
        "Reinforced concrete sections are designed to BS EN 1992-1-1 for ULS and SLS as defined in BS EN 1990. "
        "Piles and the combi wall infill are checked for N with biaxial bending with the N–M interaction of the "
        "circular section (parabola-rectangle concrete, bilinear steel), matching Oasys AdSec; each pile is "
        "divided into parts where its cage changes. The ratio below is the acting moment over the moment "
        "capacity at the same N."
    )
    rows = []
    for w in res.get("combi_walls", []):
        t = w.get("tube") or {}
        g = t.get("governing") or {}
        if t:
            zone = f", {g['zone']}" if g.get("zone") else ""
            m, mrd = (None if v is None else float(v) for v in (g.get("M_kNm"), g.get("M_Rd_kNm")))
            combo = _sheet(w["element"], g.get("combination"))
            if m is not None and mrd:
                # As the office tables: bending alone, then with the N–M (and buckling) interaction.
                rows.append(
                    [f"{w['element']} – steel tube{zone}", "N.A", round(abs(m) / mrd, 3), m, mrd, combo]
                )
            rows.append(
                [
                    f"{w['element']} – steel tube{zone} (considering interaction between moment and normal)",
                    "N.A",
                    t.get("utilisation"),
                    m,
                    mrd,
                    combo,
                ]
            )
        rows += _part_rows(f"{w['element']} – infill", w.get("infill") or {}, w["element"])
    for p in res.get("piles", []):
        rows += _part_rows(p["element"], p)
        t = (p.get("casing") or {}).get("tube") or {}
        if t:
            g = t.get("governing") or {}
            m, mrd = (None if v is None else float(v) for v in (g.get("M_kNm"), g.get("M_Rd_kNm")))
            rows.append(
                [
                    f"{p['element']} – steel casing (considering interaction between moment and normal)",
                    "N.A",
                    t.get("utilisation"),
                    m,
                    mrd,
                    _sheet(p["element"], g.get("combination")),
                ]
            )
    for b in res.get("beams", []):
        g = (b.get("bending") or {}).get("governing") or {}
        cracks = [c.get("wk") for c in (b.get("cracks") or {}).values() if c.get("wk") is not None]
        rows.append(
            [
                b["element"],
                max(cracks) if cracks else "N.A",
                (b.get("bending") or {}).get("utilisation"),
                g.get("Mv_kNm"),
                g.get("MRd_v_kNm"),
                _sheet(b["element"], g.get("combination")),
            ]
        )
    r.caption(f"Table 3-1: Summary of design results for {section.name}")
    r.table(
        [
            "Element",
            "Crack width (mm)",
            "Ultimate moment / moment capacity",
            "Acting moment (kN.m)",
            "Moment capacity (kN.m)",
            "Governing load combination",
        ],
        rows,
    )
    r.note(
        "Beams: vertical bending with the horizontal moment and N (EN 1992-1-1 5.8.9); the ratio is the check."
    )
    for s in res.get("slabs", []):
        _slab_summary(r, s)
    for w in res.get("sheet_pile_walls", []):
        _spw_summary(r, w)
    _shear_summary(r, res)
    _punching_summary(r, res)
    _steel_summary(r, res)
    failing = [
        x["element"]
        for k in ("piles", "combi_walls", "beams", "slabs")
        for x in res.get(k, [])
        if not x.get("passed")
    ]
    if failing:
        r.note("Not passing: " + ", ".join(failing) + ". See the notes of each element in the appendix.")
    for s in res.get("skipped", []):
        r.note(s)


def _sheet(element: str, combination: Any) -> Any:
    """The combination as the workbook's sheet name, e.g. Pile(1)-PT-C-Apron, as the office tables."""
    return f"{element}-{combination}" if isinstance(combination, str) and combination else combination


def _part_rows(name: str, p: dict, element: str | None = None) -> list[list[Any]]:
    """One row per station (part) of a pile: crack width, M/MRd, MEd, MRd and combination."""
    element = element or name
    stations = p.get("governing_sets") or []
    cracks = {(c["top"], c["bottom"]): c["wk_mm"] for c in (p.get("cracks") or {}).get("stations", [])}
    rows = []
    for k, st in enumerate(stations, 1):
        g = st.get("governing") or {}
        label = name if len(stations) == 1 else f"{name} – Part {k}"
        label += f" ({st['top']:g} to {st['bottom']:g} m, {st['cage']})"
        wk = cracks.get((st["top"], st["bottom"]))
        rows.append(
            [
                label,
                "N.A" if wk is None else wk,
                g.get("moment_ratio"),
                g.get("M_kNm"),
                g.get("M_Rd_kNm"),
                _sheet(element, g.get("combination")),
            ]
        )
    if not rows and p:
        g = p.get("governing") or {}
        rows.append(
            [
                name,
                (p.get("cracks") or {}).get("wk_mm"),
                p.get("utilisation"),
                g.get("M_kNm"),
                g.get("M_Rd_kNm"),
                _sheet(element, g.get("combination")),
            ]
        )
    return rows


def _strip_label(s: dict, row: dict) -> str:
    if row["strip"] == "all":  # bars over the whole deck (across the strips), or one of their zones
        return f"{s['element']} – {row['moment']} – {row['label']} – {row['face']}"
    a, b = row["station"]
    return f"{s['element']} – {row['moment']} – Station {a:g} to {b:g} – {row['strip'].capitalize()} Strip"


def _slab_summary(r: Report, s: dict) -> None:
    sd = s.get("strip_design")
    if sd and s.get("box"):
        r.image(slab_stations(sd, s["box"]), f"Figure 3-1: {s['element']} strips and stations")
    if sd:
        r.caption(
            f"Table 3-2: Summary of design results for {s['element']} ({s.get('thickness_mm', 0):g} mm)"
        )
        table = sd.get("table") or []
        r.table(
            [
                "Slab",
                "Crack width (mm)",
                "Ultimate moment / moment capacity",
                "Acting moment (kN.m/m)",
                "Moment capacity (kN.m/m)",
                "Governing load combination",
                "Bottom bars",
                "Top bars",
            ],
            [
                [
                    f"{x['moment']} – {x['label']}"
                    + ("" if x["strip"] == "all" else f" – {x['strip'].capitalize()} Strip")
                    + (" (bars set by the user)" if x.get("user_set") else ""),
                    x.get("wk_mm"),
                    x.get("ratio"),
                    x.get("M_kNm_per_m"),
                    x.get("MRd_kNm_per_m"),
                    _sheet(s["element"], x.get("combination")),
                    x["bars"].get("bottom", "–"),
                    x["bars"].get("top", "–"),
                ]
                for x in table
            ],
        )
        r.note(
            f"Stations are distances ({sd['along']}) from the {sd['from']}, on the sea side, increasing towards "
            f"the rear. Column strips are {sd['column_width_m']:g} m wide on the lines of piles, field strips "
            f"{sd['field_width_m']:g} m between them; moments are per metre, averaged across the strip, and every "
            "column (field) strip along the berth is designed together. Bars along the strips are given per "
            "station; bars along the berth are one mesh over the whole deck, with zones of additional bars "
            "only where it needs more. Each row shows its worst face; Appendix A has both."
        )
        return
    rows = []
    rest = (s.get("restraint") or {}).get("layers") or {}
    for k, lay in (s.get("layers") or {}).items():
        zones = lay.get("zones") or []
        extra = ""
        if zones:
            by_area = sorted(zones, key=lambda z: z["as_mm2_per_m"])
            extra = f"{len(zones)} zones, + {by_area[0]['label']} to + {by_area[-1]['label']} (Appendix A)"
        rows.append(
            [
                f"{s['element']} – {MOMENT.get(k, k)}",
                lay["basic"]["label"],
                extra or ("mesh only" if lay.get("mode") == "mesh_only" else "none needed"),
                (rest.get(k) or {}).get("wk"),
                lay.get("utilisation"),
            ]
        )
    r.caption(f"Table 3-2: Summary of design results for {s['element']} ({s.get('thickness_mm', 0):g} mm)")
    r.table(["Slab", "Mesh", "Additional bars", "Crack width (mm)", "Utilisation"], rows)
    r.note(
        "M11 is the moment about the slab's local 1 axis (bars along X), M22 about the 2 axis (bars along Y). "
        "Crack widths are those of restrained temperature and shrinkage on the mesh."
    )


MOMENT = {
    "bottom_x": "M11 bottom (bars along X)",
    "top_x": "M11 top (bars along X)",
    "bottom_y": "M22 bottom (bars along Y)",
    "top_y": "M22 top (bars along Y)",
}


def _shear_summary(r: Report, res: dict) -> None:
    piles = [(p["element"], p.get("shear")) for p in res.get("piles", []) if p.get("shear")]
    piles += [
        (f"{w['element']} infill", (w.get("infill") or {}).get("shear"))
        for w in res.get("combi_walls", [])
        if (w.get("infill") or {}).get("shear")
    ]
    beams = [b for b in res.get("beams", []) if (b.get("shear") or {}).get("link")]
    if not piles and not beams and not res.get("slabs"):
        return
    r.h(2, "3.2 Shear")
    if piles:
        rows = []
        for name, sh in piles:
            g = sh.get("governing") or {}
            links = ", ".join(dict.fromkeys(z["link"] for z in sh.get("zones", [])))
            rows.append([name, g.get("V_kN"), g.get("N_kN"), links, sh.get("utilisation")])
        r.caption("Table 3-3: Shear in the piles")
        r.table(["Pile", "Q (kN)", "Nu (kN)", "As (links)", "Utilisation"], rows)
    if beams:
        rows = []
        for b in beams:
            sh = b["shear"]
            g = sh.get("governing") or {}
            link = sh["link"]
            rows.append(
                [
                    b["element"],
                    g.get("V_kN"),
                    g.get("N_kN"),
                    f"Ø{link['phi']:g} @ {link['spacing_mm']:g}",
                    link.get("legs"),
                    sh.get("utilisation"),
                ]
            )
        r.caption("Table 3-4: Shear in the beams")
        r.table(["Beam", "Q (kN)", "Nu (kN)", "As", "Number of legs", "Utilisation"], rows)
    for s in res.get("slabs", []):
        sh = s.get("shear") or {}
        if not sh:
            continue
        g = sh.get("governing") or {}
        r.p(
            f"{s['element']}: one-way shear per metre, {sh.get('method', '')}. Largest {_fmt(g.get('V_kN_per_m'))} kN/m "
            f"({g.get('combination')}, X {g.get('x')}, Y {g.get('y')}) against VRd,c {_fmt(g.get('VRd_c_kN_per_m'))} kN/m; "
            f"links in {sh.get('cells_needing_links', 0)} cells"
            + (f", heaviest {sh['heaviest']['label']}." if sh.get("heaviest") else ".")
        )


def _punching_summary(r: Report, res: dict) -> None:
    rows = []
    for s in res.get("slabs", []):
        for q in s.get("punching") or []:
            face = q["vEd_face_MPa"] / q["vRd_max_MPa"]
            if face > 1:
                links = f"FAILS at the pile face: vEd,0 {q['vEd_face_MPa']:.2f} > vRd,max {q['vRd_max_MPa']:.2f} MPa"
            elif q.get("kmax_ratio", 0) > 1 and q.get("needs_reinforcement"):
                links = "FAILS: over kmax · vRd,c, links cannot carry it"
            elif q.get("perimeters"):
                links = f"{q['perimeters']} perimeters @ {q['radial_spacing_mm']} mm, {q['asw_mm2_per_perimeter']} mm² each"
            else:
                links = "NO NEED FOR R.F.T"
            rows.append(
                [
                    f"{q['pile']} ({q['x']:g}, {q['y']:g})",
                    q.get("utilisation"),
                    q.get("utilisation_with_links")
                    if q.get("perimeters")
                    else (round(face, 3) if face > 1 else "-"),
                    links,
                    q.get("combination"),
                ]
            )
    if rows:
        r.h(2, "3.3 Punching")
        r.caption("Table 3-5: Punching of the slab over the piles (EN 1992-1-1 6.4)")
        r.table(
            [
                "Pile (X, Y)",
                "Punching utilisation, concrete only",
                "Punching utilisation, concrete and reinforcement",
                "Punching reinforcement",
                "Governing load combination",
            ],
            rows,
        )


def overall_ratio(steel: dict) -> float | None:
    """The main bars' volume over the element's concrete, in %; worked out for older results."""
    if steel.get("ratio_pct") is not None:
        return steel["ratio_pct"]
    if steel.get("longitudinal_kg") and steel.get("concrete_m3"):
        return round(100 * steel["longitudinal_kg"] / STEEL_DENSITY / steel["concrete_m3"], 2)
    return None


def _steel_summary(r: Report, res: dict) -> None:
    rows = []
    for p in res.get("piles", []):
        st = p.get("steel") or {}
        rows.append(
            [
                p["element"],
                overall_ratio(st),
                st.get("kg_per_m3", p.get("steel_ratio_kg_m3")),
                st.get("element_total_t"),
            ]
        )
    for w in res.get("combi_walls", []):
        st = (w.get("infill") or {}).get("steel") or {}
        rows.append(
            [f"{w['element']} infill", overall_ratio(st), st.get("kg_per_m3"), st.get("element_total_t")]
        )
    for b in res.get("beams", []):
        st = b.get("steel") or {}
        rows.append(
            [
                b["element"],
                overall_ratio(st),
                st.get("kg_per_m3"),
                st.get("element_total_t", st.get("total_t")),
            ]
        )
    for s in res.get("slabs", []):
        st = s.get("steel") or {}
        rows.append([s["element"], overall_ratio(st), st.get("kg_per_m3"), st.get("total_t")])
    if rows:
        r.h(2, "3.4 Reinforcement quantities")
        r.caption("Table 3-6: Reinforcement per element")
        r.table(["Element", "Overall ρ, main bars (%)", "kg/m³ incl. links", "Total (t)"], rows)
        r.note(
            "Overall ρ is the main bars' volume over the whole element's concrete, laps included; the "
            "ratio at the pile head, where the cage is heaviest, is in each pile's calculation."
        )


def _sets(r: Report, sets: list[dict], title: str) -> None:
    for st in sets or []:
        head = f"{title}, {st.get('top', '')} to {st.get('bottom', '')} m, {st.get('cage', '')}"
        r.h(3, head)
        for key, name in (("uls", "ULS"), ("qp", "QP")):
            rows = [
                [
                    name,
                    x.get("case"),
                    x.get("N_kN"),
                    x.get("M2_kNm"),
                    x.get("M3_kNm"),
                    x.get("combination"),
                    x.get("z"),
                    x.get("utilisation"),
                ]
                for x in st.get(key) or []
            ]
            r.table(["", "Case", "N kN", "M2 kNm", "M3 kNm", "Combination", "z m", "Utilisation"], rows)


def _pile_body(r: Report, p: dict) -> None:
    a = p.get("arrangement") or {}
    sec = p.get("section") or {}
    g = p.get("governing") or {}
    r.kv(
        [
            ("Diameter", f"{sec.get('diameter_mm', 0):g} mm"),
            (
                "Cover to links / link Ø",
                f"{sec.get('cover_mm', 0):g} / {sec.get('link_diameter_mm', 0):g} mm",
            ),
            ("Top / toe level", f"{sec.get('head_level_m')} / {sec.get('toe_level_m')} m"),
            (
                "Cage at the head",
                f"{a.get('label')} (set by the user, checked)" if p.get("user_set") else a.get("label"),
            ),
            ("Steel area", f"{a.get('area_mm2', 0):,} mm² ({p.get('reinforcement_ratio_pct', 0):.2f}%)"),
            ("N–M utilisation", p.get("utilisation")),
            ("Result", _ok(p.get("passed"))),
        ]
    )
    if a.get("rings"):
        r.table(
            ["Row", "Bars", "Radius mm", "Clear spacing mm"],
            [
                [i + 1, f"{g_['count']}Ø{g_['diameter']}", g_["radius"], g_["clear_spacing_mm"]]
                for i, g_ in enumerate(a["rings"])
            ],
        )
    if g:
        r.p(
            f"Governing N–M: {g.get('combination')} at z {g.get('z')} m, NEd = {_fmt(g.get('N_kN'))} kN, "
            f"MEd = {_fmt(g.get('M_kNm'))} kNm, MRd at this N = {_fmt(g.get('M_Rd_kNm'))} kNm."
        )
    runs = (p.get("curtailment") or {}).get("runs") or []
    if runs:
        r.h(3, "Reinforcement down the pile")
        r.table(
            ["From m", "To m", "Cage", "Bar lengths m", "Above head m", "Lap below m", "Utilisation"],
            [
                [
                    x["top"],
                    x["bottom"],
                    x["cage"]["label"],
                    ", ".join(f"{v:g}" for v in x.get("bar_lengths_m", [])),
                    ", ".join(f"{v:g}" for v in x.get("above_head_m", [])) or "-",
                    ", ".join(f"{v:g}" for v in x.get("lap_below_m", [])),
                    x.get("utilisation"),
                ]
                for x in runs
            ],
        )
    sh = p.get("shear")
    if sh:
        r.h(3, "Shear")
        r.p(sh.get("method", ""))
        r.table(
            ["From m", "To m", "Links", "Reason"],
            [[z["top"], z["bottom"], z["link"], z.get("reason")] for z in sh.get("zones", [])],
        )
        sg = sh.get("governing") or {}
        if sg:
            r.p(
                f"Governing shear: {sg.get('combination')} at z {sg.get('z')} m, VEd = {_fmt(sg.get('V_kN'))} kN, "
                f"N = {_fmt(sg.get('N_kN'))} kN, VRd,c = {_fmt(sg.get('VRd_c_kN'))} kN, VRd,max = {_fmt(sg.get('VRd_max_kN'))} kN; "
                f"utilisation {_fmt(sh.get('utilisation'))}."
            )
    c = p.get("cracks")
    if c and c.get("wk_mm") is not None:
        cg = c.get("governing") or {}
        r.h(3, "Crack width (QP)")
        r.p(
            f"wk = {c['wk_mm']:.3f} mm against {c['limit_mm']:g} mm ({_ok(c.get('passed'))}), at z {cg.get('z')} m, "
            f"{cg.get('combination')}: N = {_fmt(cg.get('N_kN'))} kN, M = {_fmt(cg.get('M_kNm'))} kNm, "
            f"σs = {_fmt(cg.get('sigma_s_MPa'))} MPa, sr,max = {_fmt(cg.get('sr_max_mm'))} mm (EN 1992-1-1 7.3.4)."
        )
    if c and c.get("casing"):
        r.note(c["casing"])
    cas = p.get("casing")
    if cas and cas.get("tube"):
        t = cas["tube"]
        sc = t.get("section") or {}
        g = t.get("governing") or {}
        r.h(3, "Steel casing (structural)")
        r.p(cas["note"])
        r.kv(
            [
                (
                    "Casing",
                    f"Ø{_fmt(sc.get('diameter_mm'))} × {_fmt(sc.get('thickness_mm'))} mm {sc.get('grade', '')}",
                ),
                (
                    "Corroded",
                    f"Ø{_fmt(sc.get('corroded_diameter_mm'))} × {_fmt(sc.get('corroded_thickness_mm'))} mm",
                ),
                *[(k.replace("_", " "), v) for k, v in (t.get("resistances") or {}).items()],
                (
                    "Governing",
                    f"{g.get('combination')}, z {g.get('z')} m: N = {_fmt(g.get('N_kN'))} kN, "
                    f"M = {_fmt(g.get('M_kNm'))} kNm, V = {_fmt(g.get('V_kN'))} kN ({g.get('check')})",
                ),
                ("Utilisation", t.get("utilisation")),
                ("Result", _ok(t.get("passed"))),
            ]
        )
    con = p.get("connection")
    if con:
        r.h(3, "Casing connection")
        r.p(
            f"{con.get('bottom')} to {con.get('top')} m, welded {con.get('welded') or 'none'} with {con.get('cage')}: "
            f"utilisation {_fmt(con.get('utilisation'))}."
        )
    st = p.get("steel") or {}
    if st:
        r.kv(
            [
                ("Longitudinal bars (with laps)", f"{st.get('longitudinal_kg', 0):,.0f} kg"),
                (
                    "Links",
                    f"{st.get('links_kg', 0):,.0f} kg"
                    + (
                        f" (incl. {sh_['inner_links_kg']:,.0f} kg in {sh_['inner_rings']} inner ring(s))"
                        if (sh_ := p.get("shear") or {}).get("inner_rings")
                        else ""
                    ),
                ),
                ("Steel ratio", f"{st.get('kg_per_m3', 0):.0f} kg/m³"),
                ("Piles of this type", p.get("count")),
                (
                    "Steel for all of them",
                    f"{st['element_total_t']:.1f} t" if st.get("element_total_t") is not None else None,
                ),
            ]
        )
    for n in p.get("notes", []):
        r.note(n)
    _sets(r, p.get("governing_sets"), "Governing sets")


def _pile(r: Report, p: dict) -> None:
    r.h(1, f"{p['element']}: pile")
    r.p(
        "N–M interaction to EN 1992-1-1 (parabola-rectangle concrete, bilinear steel, 6.1 strain limits), "
        "the worst bar orientation, utilisation along the ray from the origin to each load."
    )
    _pile_body(r, p)


def _combi(r: Report, w: dict) -> None:
    r.h(1, f"{w['element']}: combi wall")
    r.kv(
        [
            ("King piles", w.get("count")),
            ("Infill bottom level", f"{w.get('infill_bottom_level')} m"),
            ("Share carried by the steel (E·I)", f"{w.get('steel_share', 0) * 100:.0f}%"),
            (
                "Tube check",
                "office sheets: elastic, class 4 effective properties"
                if (w.get("tube") or {}).get("method") == "office"
                else "EN 1993: plastic where filled, shell buckling where empty",
            ),
            ("Utilisation", w.get("utilisation")),
            ("Result", _ok(w.get("passed"))),
        ]
    )
    t = w.get("tube")
    if t:
        s = t["section"]
        r.h(2, "Steel tube")
        r.kv(
            [
                ("Tube", f"Ø{s['diameter_mm']:g} × {s['thickness_mm']:g} mm {s.get('grade', '')}"),
                (
                    "Corroded",
                    f"Ø{s['corroded_diameter_mm']:g} × {s['corroded_thickness_mm']:g} mm (loss {s['corrosion_mm']:g} mm)",
                ),
                ("Class unfilled", s.get("class_unfilled")),
            ]
        )
        r.kv([(k.replace("_", " "), v) for k, v in (t.get("resistances") or {}).items()])
        zones = t.get("zones") or []
        if len(zones) > 1 or (zones and zones[0].get("inside_mm")):
            r.caption(f"Corrosion zones of the {w['element']} tube")
            r.table(
                [
                    "From (m)",
                    "To (m)",
                    "Outside (mm)",
                    "Inside (mm)",
                    "t (mm)",
                    "Class",
                    "Aeff (mm²)",
                    "Meff (kN.m)",
                ],
                [
                    [
                        "top" if z.get("top") is None else z["top"],
                        "toe" if z.get("bottom") is None else z["bottom"],
                        z.get("outside_mm"),
                        z.get("inside_mm"),
                        z.get("t_mm"),
                        z.get("class"),
                        z.get("A_eff_mm2"),
                        z.get("M_eff_kNm"),
                    ]
                    for z in zones
                ],
            )
        col = t.get("column")
        if col:
            r.h(3, "Column buckling (composite, EN 1994-1-1 6.7.3)")
            r.kv(
                [
                    ("Length / buckling length", f"{col['length_m']} m / {col['buckling_length_m']} m"),
                    ("EI,eff = EaIa + 0.6·Ecm·Ic (average)", f"{_fmt(col['EI_eff_kNm2'])} kN.m²"),
                    ("Npl,Rk (average)", f"{_fmt(col['N_pl_Rk_kN'])} kN"),
                    ("Ncr", f"{_fmt(col['N_cr_kN'])} kN"),
                    ("λ̄ / curve / χ", f"{col['slenderness']} / {col['curve']} / {col['chi']}"),
                    ("Nb,Rd", f"{_fmt(col['N_b_Rd_kN'])} kN"),
                    ("NEd/Nb,Rd + kyy·MEd/Meff,Rd (Cmy 0.9)", col.get("utilisation")),
                ]
            )
        if t.get("buckling"):
            r.kv([(f"Shell buckling {k.replace('_', ' ')}", v) for k, v in t["buckling"].items()])
        g = t.get("governing") or {}
        r.p(
            f"Governing: {g.get('combination')} at z {g.get('z')} m ({g.get('zone')}), N = {_fmt(g.get('N_kN'))} kN; "
            f"utilisation {_fmt(t.get('utilisation'))}."
        )
        for n in t.get("notes", []):
            r.note(n)
        gs = t.get("governing_sets") or {}
        if gs.get("rows"):
            cols = list(gs.get("columns", {}).keys()) or ["N", "M2", "M3", "Q1", "Q2"]
            r.table(
                ["Case", *cols, "Combination", "z m"],
                [
                    [x.get("case"), *[x.get(c) for c in cols], x.get("combination"), x.get("z")]
                    for x in gs["rows"]
                ],
            )
    inf = w.get("infill")
    if inf:
        r.h(2, "Concrete infill")
        _pile_body(r, inf)
    for n in w.get("notes", []):
        r.note(n)


def _beam(r: Report, b: dict) -> None:
    r.h(1, f"{b['element']}: {KIND.get(b.get('kind'), 'beam').lower()}")
    c = b.get("cage") or {}
    r.kv(
        [
            (
                "Section",
                f"{b.get('width_mm', 0):g} × {b.get('depth_mm', 0):g} mm, {b.get('concrete')}, cover {b.get('cover_mm', 0):g} mm",
            ),
            ("Along", f"{b.get('along')} from {b.get('start_m')} to {b.get('end_m')} m"),
            (
                "Supports",
                ", ".join(f"{s.get('element')} at {s.get('s')} m" for s in b.get("supports", [])) or "none",
            ),
            ("Cage", c.get("label")),
            ("Steel area", f"{c.get('area_mm2', 0):,} mm² ({c.get('ratio_pct', 0):.2f}%)"),
            ("Utilisation (all checks)", b.get("utilisation")),
            ("Result", _ok(b.get("passed"))),
        ]
    )
    bend = b.get("bending") or {}
    g = bend.get("governing") or {}
    if g:
        r.h(3, "N with biaxial bending (EN 1992-1-1 5.8.9)")
        r.p(
            f"{g.get('combination')} at {g.get('s')} m: N = {_fmt(g.get('N_kN'))} kN, Mv = {_fmt(g.get('Mv_kNm'))} kNm "
            f"(MRd {_fmt(g.get('MRd_v_kNm'))}), Mh = {_fmt(g.get('Mh_kNm'))} kNm (MRd {_fmt(g.get('MRd_h_kNm'))}), "
            f"a = {_fmt(g.get('a'))}; utilisation {_fmt(bend.get('utilisation'))}."
        )
    ex = bend.get("extremes") or {}
    r.table(["Action", "Max", "Min"], [[k, v.get("max"), v.get("min")] for k, v in ex.items()])
    faces = b.get("faces") or []
    if faces:
        r.h(3, "What sets each face (mm² each check needs)")

        def need(f: dict, k: str) -> Any:
            v = (f.get("needs_mm2") or {}).get(k, "–")
            return "below minimum" if v == 0 else "more than any bars" if v is None else v

        keys = ["minimum", "bending", "crack", "restraint"] + (
            ["truss"] if any("truss" in (f.get("needs_mm2") or {}) for f in faces) else []
        )
        r.table(
            ["Face", "Minimum", "Bending (Plaxis)", "QP crack (Plaxis)", "Restraint", "Truss tie"][
                : len(keys) + 1
            ]
            + ["From Plaxis alone", "Final", "Governed by"],
            [
                [
                    f["face"],
                    *[need(f, k) for k in keys],
                    f.get("plaxis"),
                    f.get("final"),
                    f.get("governed_by"),
                ]
                for f in faces
            ],
        )
    cr = b.get("cracks") or {}
    rs = (b.get("restraint") or {}).get("faces") or {}
    rows = [
        ["QP load", f, x.get("wk"), x.get("limit"), x.get("sigma_s"), x.get("sr_max"), _ok(x.get("passed"))]
        for f, x in cr.items()
    ]
    rows += [
        ["Restraint", f, x.get("wk"), x.get("limit"), None, x.get("sr_max"), _ok(x.get("passed"))]
        for f, x in rs.items()
    ]
    if rows:
        r.h(3, "Crack widths")
        r.table(["Check", "Face", "wk mm", "Limit mm", "σs MPa", "sr,max mm", "Result"], rows)
    sh = b.get("shear") or {}
    if sh.get("link"):
        r.h(3, "Links")
        sg = sh.get("governing") or {}
        r.p(
            f"{sh['link']['label']} for the whole beam ({sh.get('method', '')}). Governing {sg.get('combination')} at "
            f"{sg.get('s')} m: V = {_fmt(sg.get('V_kN'))} kN, T = {_fmt(sg.get('T_kNm'))} kNm, N = {_fmt(sg.get('N_kN'))} kN; "
            f"utilisation {_fmt(sh.get('utilisation'))}. Torsion longitudinal steel {_fmt(sh.get('torsion_long_steel_mm2'))} mm²."
        )
    tr = b.get("transverse") or {}
    if tr.get("top"):
        r.h(3, "Transverse bars per metre")
        r.p(
            f"Top {tr['top']['label']}, bottom {tr['bottom']['label']}; utilisation {_fmt(tr.get('utilisation'))}."
        )
    st = b.get("steel") or {}
    r.kv(
        [
            ("Steel ratio", f"{st.get('kg_per_m3', 0):.0f} kg/m³"),
            ("Per metre", f"{st.get('kg_per_m', 0):.0f} kg/m"),
        ]
    )
    bo = b.get("bollard")
    if bo:
        r.h(3, "Bollard tie bars")
        r.p(bo["method"])
        r.table(
            ["Tie bars", "Plan angle (°)", "T_Rd (kN)", "Square to the quay (kN)", "Lap l0 needed (mm)"],
            [
                [t["bars"], t["angle_deg"], t["T_Rd_kN"], t["normal_kN"], lap["l0_mm"]]
                for t, lap in zip(bo["ties"], bo["laps"], strict=False)
            ],
        )
        r.kv(
            [
                (
                    "Bollard",
                    f"{bo['capacity_t']:g} t × {bo['load_factor']:g}: F_Ed = {_fmt(bo['F_Ed_kN'])} kN",
                ),
                ("Ties", f"{_fmt(bo['R_kN'])} kN, utilisation {bo['tie_utilisation']}"),
                ("Lap given", f"{bo['lap_length_mm']:g} mm, utilisation {bo['lap_utilisation']}"),
                ("Result", _ok(bo.get("passed"))),
            ]
        )
    tr = b.get("truss")
    if tr and tr.get("cases"):
        r.h(3, "Truss model between king piles")
        r.p(tr["method"])
        r.kv(
            [
                (
                    "King pile spacing",
                    f"{tr['spacing_m']:g} m ({'input' if tr['spacing_from'] == 'input' else 'from the workbook'})",
                ),
                ("Lever arm z", f"{_fmt(tr['lever_arm_mm'])} mm"),
                (
                    "θ = atan(z / (s/2))",
                    f"{tr['theta_deg']:g}°: F = {tr['strut_factor']:g} P, T = {tr['tie_factor']:g} P",
                ),
                ("Bottom bars", f"{_fmt(tr['As_provided_mm2'])} mm² at {tr['working_stress_MPa']:g} MPa"),
            ]
        )
        r.table(
            ["Case", "Slab (mm)", "P beam (kN)", "P slab (kN)", "P (kN)", "T (kN)", "As req (mm²)", "Ratio"],
            [
                [c["case"]]
                + [
                    float(c[k])
                    for k in ("slab_thickness_mm", "P_beam_kN", "P_slab_kN", "P_kN", "T_kN", "As_req_mm2")
                ]
                + [c["utilisation"]]
                for c in tr["cases"]
            ],
        )
        r.kv([("Result", _ok(tr.get("passed")))])
    elif tr:
        r.note(tr.get("note", ""))
    for n in b.get("notes", []):
        r.note(n)
    _sets(r, b.get("governing_sets"), "Governing sets")


LAYER = {"bottom_x": "bottom X", "bottom_y": "bottom Y", "top_x": "top X", "top_y": "top Y"}


def _slab(r: Report, s: dict) -> None:
    r.h(1, f"{s['element']}: slab")
    st = s.get("steel") or {}
    r.kv(
        [
            ("Thickness", f"{s.get('thickness_mm', 0):g} mm, {s.get('concrete')}"),
            ("Covers top / bottom", f"{s.get('cover_top_mm', 0):g} / {s.get('cover_bottom_mm', 0):g} mm"),
            ("Layout", "column and field strips" if s.get("strips") == "column_and_field" else "uniform"),
            (
                "Steel",
                f"{st.get('kg_per_m3', 0):.0f} kg/m³, {st.get('kg_per_m2', 0):.1f} kg/m², {st.get('total_t', 0):.1f} t (links not included)",
            ),
            ("Utilisation", s.get("utilisation")),
            ("Result", _ok(s.get("passed"))),
        ]
    )
    sd = s.get("strip_design")
    if sd:
        r.h(3, "Column and field strips")
        r.p(
            f"Strips along {sd['along']} from the {sd['from']}: column strips {sd['column_width_m']:g} m wide on the "
            f"lines of piles at {sd['along'] == 'X' and 'Y' or 'X'} = "
            + ", ".join(f"{v:g}" for v in sd["lines"])
            + f" m, field strips {sd['field_width_m']:g} m between them. Stations: "
            + ", ".join(f"{v:g}" for v in sd["stations"])
            + " m. At every cut along a strip the Wood–Armer moment and N are averaged across its width; the "
            "worst cut of any column (field) strip in a station sets that station's bars, and the QP crack width "
            "is checked the same way. MRd with the tension bars only, rectangular block 0.8x."
        )
        r.table(
            [
                "Strip",
                "Face",
                "MEd kN.m/m",
                "NEd kN/m",
                "Combination",
                "Bars",
                "MRd kN.m/m",
                "M/MRd",
                "wk mm",
            ],
            [
                [
                    _strip_label(s, x),
                    x["face"],
                    x["M_kNm_per_m"],
                    x["N_kN_per_m"],
                    x["combination"],
                    x["bars"],
                    x["MRd_kNm_per_m"],
                    x["ratio"],
                    x["wk_mm"],
                ]
                for x in sd["rows"]
            ],
        )
    r.h(3, "Bars per metre")
    r.p(
        "Wood–Armer moments with the in-plane N; a mesh everywhere and additional bars between its bars where needed."
    )
    rows = []
    for k, lay in (s.get("layers") or {}).items():
        rows.append(
            [
                LAYER.get(k, k),
                f"{lay['basic']['label']} ({lay['basic']['as_mm2_per_m']:,} mm²/m)",
                "mesh only" if lay.get("mode") == "mesh_only" else f"{len(lay.get('zones', []))} zones",
                lay.get("utilisation"),
                lay.get("d_mm"),
            ]
        )
    r.table(["Layer", "Mesh", "Additional bars", "Utilisation", "d mm"], rows)
    for k, lay in (s.get("layers") or {}).items():
        if lay.get("zones"):
            r.h(3, f"Additional bars, {LAYER.get(k, k)}")
            r.table(
                ["X from", "X to", "Y from", "Y to", "Additional", "Total mm²/m"],
                [
                    [z["x"][0], z["x"][1], z["y"][0], z["y"][1], z["label"], z["as_mm2_per_m"]]
                    for z in lay["zones"]
                ],
            )
    punch = s.get("punching") or []
    if punch:
        r.h(3, "Punching (EN 1992-1-1 6.4)")
        r.table(
            [
                "Pile",
                "X, Y",
                "h mm",
                "VEd kN",
                "β",
                "vEd / vRd,c MPa",
                "Face / vRd,max MPa",
                "Links",
                "Result",
            ],
            [
                [
                    q["pile"],
                    f"{q['x']:g}, {q['y']:g}",
                    q.get("thickness_mm"),
                    q.get("V_kN"),
                    q.get("beta"),
                    f"{q['vEd_MPa']:.3f} / {q['vRd_c_MPa']:.3f}",
                    f"{q['vEd_face_MPa']:.2f} / {q['vRd_max_MPa']:.2f}",
                    f"{q['perimeters']} perimeters @ {q['radial_spacing_mm']} mm, {q['asw_mm2_per_perimeter']} mm² each"
                    if q.get("perimeters")
                    else ("needed" if q.get("needs_reinforcement") else "none"),
                    _ok(q.get("passed")),
                ]
                for q in punch
            ],
        )
    sh = s.get("shear") or {}
    if sh:
        r.h(3, "Shear per metre")
        g = sh.get("governing") or {}
        r.p(
            f"{sh.get('method', '')}. Largest: {g.get('combination')} at X {g.get('x')}, Y {g.get('y')}: v = {_fmt(g.get('V_kN_per_m'))} kN/m, "
            f"VRd,c = {_fmt(g.get('VRd_c_kN_per_m'))} kN/m, VRd,max = {_fmt(g.get('VRd_max_kN_per_m'))} kN/m. "
            f"Links in {sh.get('cells_needing_links', 0)} cells"
            + (f", heaviest {sh['heaviest']['label']}." if sh.get("heaviest") else ".")
        )
    rest = s.get("restraint") or {}
    if rest.get("layers"):
        r.h(3, "Temperature and shrinkage restraint")
        r.table(
            ["Layer", "wk mm", "Limit mm", "εr µε", "sr,max mm", "Result"],
            [
                [
                    LAYER.get(k, k),
                    x.get("wk"),
                    x.get("limit"),
                    x.get("eps_r"),
                    x.get("sr_max"),
                    _ok(x.get("passed")),
                ]
                for k, x in rest["layers"].items()
            ],
        )
        r.p(f"{rest.get('length_m')} m between joints, R = {rest.get('R')} ({rest.get('R_from')}).")
    for n in s.get("notes", []):
        r.note(n)


def _spw_summary(r: Report, w: dict) -> None:
    d = w.get("design")
    if not d or d.get("error"):
        r.p(f"{w['element']}: {d['error'] if d else 'no sheet pile design (no results)'}")
        return
    t = d["check_titles"]
    g = d["designed"]["governing"]
    r.p(
        f"{w['element']}: {d['section']}, fy {d['steel']['fy']:g} MPa, checked to EN 1993-5 as ArcelorMittal "
        f"Durability 4.2.1 at every Plaxis result: Uf = {d['uf']:.2f} ({t[g['governs']].lower()}, "
        f"{g['combination']}, z {g['z']:.2f} m), {'passes' if d['ok'] else 'FAILS'}."
    )
    if d["adjusted"]:
        p = d["as_plaxis"]
        left = "; ".join(
            f"{' and '.join(x for x, on in (('N', i['ignore_n']), ('Q', i['ignore_q'])) if on)} left out in "
            f"{i['combination'].lower() if i['combination'] == 'All combinations' else i['combination']}"
            for i in d["ignored"]
        )
        r.note(
            f"{w['element']}: {left}, as the Plaxis values there are not taken as sheet pile actions. With every "
            f"Plaxis action Uf = {p['uf']:.2f} ({t[p['governing']['governs']].lower()})."
        )


def _spw(r: Report, w: dict) -> None:
    r.h(1, f"{w['element']}: sheet pile wall")
    d = w.get("design")
    if d and not d.get("error"):
        t = d["check_titles"]
        pr, st = d["properties"], d["settings"]
        r.p(
            f"{d['section']} ({pr['h']:g} mm high, tf {pr['tf']:g} mm, tw {pr['tw']:g} mm; A {pr['area']:g} cm²/m, "
            f"I {pr['inertia']:g} cm⁴/m, Wel {pr['wel']:g} cm³/m, Wpl {pr['wpl']:g} cm³/m), fy {d['steel']['fy']:g} "
            f"MPa, γM0 {st['gamma_m0']:g}, γM1 {st['gamma_m1']:g}, buckling length {st['buckling_length']:g} m"
            f"{'' if st['buckling_length_given'] else ' (assumed)'}. Flange b {pr['flange']:g} mm"
            f"{'' if pr['flange_given'] else ' (estimated)'}, web angle {pr['angle']:g}°"
            f"{'' if pr['angle_given'] else ' (estimated)'}. Checked at {d['points']} Plaxis points with M = |M_11| "
            f"+ |N| e, V = |{st['shear']}|, N = −N_1 (compression +); corrosion (front + back) taken off every plate."
        )
        r.table(
            ["Zone", "Down to m", "Front mm", "Back mm", "Total mm"],
            [[z["zone"], z["bottom"], z["front"], z["back"], z["total"]] for z in d["zones"]],
        )
        for key, title in (("designed", "Results by zone" + (" (as designed)" if d["adjusted"] else "")),) + (
            (("as_plaxis", "Results by zone with every Plaxis action"),) if d["adjusted"] else ()
        ):
            res = d[key]
            r.h(2, title)
            r.table(
                ["Zone", "z m", "Combination", "M kNm/m", "V kN/m", "N kN/m", "Loss mm", "Class"]
                + [t[c] for c in t]
                + ["Uf"],
                [
                    [
                        z["zone"],
                        z["z"],
                        z["combination"],
                        z["M"],
                        z["V"],
                        z["N"],
                        z["loss"],
                        z["values"]["class"],
                    ]
                    + [("–" if z["checks"][c] is None else z["checks"][c]) for c in t]
                    + [z["uf"]]
                    for z in res["zones"]
                ],
            )
        g = d["designed"]["governing"]
        v = g["values"]
        r.h(2, "Governing point")
        r.bullets(
            [
                f"{g['combination']}, z {g['z']:.2f} m, corrosion {g['loss']:g} mm: tf {v['tf']:.2f} mm, tw "
                f"{v['tw']:.2f} mm, A {v['area']:.1f} cm²/m, I {v['inertia']:.0f} cm⁴/m, Wel {v['wel']:.0f}, Wpl "
                f"{v['wpl']:.0f} cm³/m; (b / tf) / ε = {v['slender']:.1f}, class {v['class']:g}"
                + (f" taken as 3 with fy,red {v['fy_used']:.1f} MPa" if v["class"] == 4 else ""),
                f"MEd {g['M']:.1f} kNm/m, Mc,Rd {v['Mc']:.1f} kNm/m; VEd {g['V']:.1f} kN/m, Vpl,Rd {v['Vpl']:.1f} kN/m"
                + (f", Vb,Rd {v['Vb']:.1f} kN/m" if v.get("Vb") is not None else "")
                + f"; NEd {g['N']:.1f} kN/m, Npl,Rd {v['Npl']:.0f} kN/m, Ncr {v['Ncr']:.0f} kN/m, χ {v['chi']:.3f}",
                f"Uf = {g['uf']:.3f} ({t[g['governs']].lower()})",
            ]
        )
        for n in [*(w.get("notes") or []), *d["notes"]]:
            r.note(n)
    r.p("Governing straining actions (Plaxis sign, multipliers applied), for Durability:")
    gs = w.get("governing_sets") or {}
    cols = list((gs.get("columns") or {}).keys())
    r.table(
        ["Case", *cols, "Combination", "z m"],
        [
            [x.get("case"), *[x.get(c) for c in cols], x.get("combination"), x.get("z")]
            for x in gs.get("rows", [])
        ],
    )


# --- Renderers --------------------------------------------------------------------------------------


def to_xlsx(rep: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold, head = Font(bold=True), PatternFill("solid", fgColor="E8EEF4")
    row = 1

    def sheet_for(title: str):
        name = "".join(ch for ch in title.split(":")[0] if ch not in "[]:*?/\\")[:31] or "Element"
        base, k = name, 2
        while name in wb.sheetnames:
            name = f"{base[:28]} {k}"
            k += 1
        return wb.create_sheet(name)

    ws.cell(row, 1, rep.title).font = Font(bold=True, size=14)
    ws.cell(row + 1, 1, rep.subtitle)
    row += 3
    for b in rep.blocks:
        if b.kind == "h1" and b.text.startswith("Appendix"):
            row += 1
        elif b.kind == "h1" and not b.text[:1].isdigit():
            ws = sheet_for(b.text)
            row = 1
        if b.kind.startswith("h"):
            ws.cell(row, 1, b.text).font = Font(bold=True, size={"h1": 13, "h2": 12}.get(b.kind, 11))
            row += 1
        elif b.kind == "caption":
            ws.cell(row, 1, b.text).font = Font(bold=True, italic=True)
            row += 1
        elif b.kind == "image":
            from openpyxl.drawing.image import Image as XlImage

            pic = XlImage(io.BytesIO(b.image))
            pic.width, pic.height = pic.width * 0.5, pic.height * 0.5
            ws.add_image(pic, f"A{row}")
            row += int(pic.height / 20) + 2
        elif b.kind == "bullets":
            for x in b.rows:
                ws.cell(row, 1, "• " + x[0])
                row += 1
        elif b.kind in ("p", "note"):
            c = ws.cell(row, 1, b.text)
            c.alignment = Alignment(wrap_text=False)
            if b.kind == "note":
                c.font = Font(italic=True, color="666666")
            row += 1
        elif b.kind == "kv":
            for k, v in b.rows:
                ws.cell(row, 1, k).font = bold
                ws.cell(row, 2, v)
                row += 1
            row += 1
        elif b.kind == "table":
            for j, h in enumerate(b.headers, 1):
                c = ws.cell(row, j, h)
                c.font, c.fill = bold, head
            row += 1
            for rr in b.rows:
                for j, v in enumerate(rr, 1):
                    ws.cell(row, j, _number(v))
                row += 1
            row += 1
    for s in wb.worksheets:
        s.column_dimensions["A"].width = 30
        for col in "BCDEFGHI":
            s.column_dimensions[col].width = 18
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _number(v: str) -> Any:
    try:
        return (
            float(v.replace(",", ""))
            if isinstance(v, str) and v.replace(",", "").replace(".", "", 1).lstrip("-").isdigit()
            else v
        )
    except ValueError:
        return v


def to_docx(rep: Report) -> bytes:
    """A4 portrait Word report: title, headings, captions above tables, grid tables."""
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))
    for side in ("top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.0))
    doc.styles["Normal"].font.size = Pt(10)
    footer = sec.footer.paragraphs[0]
    footer.text = f"{rep.title} · {rep.subtitle}"
    footer.runs[0].font.size = Pt(7)
    doc.add_heading(rep.title, 0)
    doc.add_paragraph(rep.subtitle)
    for b in rep.blocks:
        if b.kind.startswith("h"):
            doc.add_heading(b.text, int(b.kind[1]))
        elif b.kind == "p":
            doc.add_paragraph(b.text)
        elif b.kind == "image":
            doc.add_picture(io.BytesIO(b.image), width=Cm(16))
        elif b.kind == "caption":
            doc.add_paragraph(b.text, style="Caption")
        elif b.kind == "bullets":
            for x in b.rows:
                doc.add_paragraph(x[0], style="List Bullet")
        elif b.kind == "note":
            run = doc.add_paragraph().add_run(b.text)
            run.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        elif b.kind in ("kv", "table"):
            headers = b.headers if b.kind == "table" else []
            t = doc.add_table(rows=0, cols=len(b.rows[0]) if b.rows else len(headers))
            t.style = "Table Grid"
            small = Pt(8 if len(headers) > 5 else 9)
            if headers:
                cells = t.add_row().cells
                for c, h in zip(cells, headers, strict=True):
                    c.text = h
                    run = c.paragraphs[0].runs[0]
                    run.bold, run.font.size = True, small
            for rr in b.rows:
                cells = t.add_row().cells
                for j, (c, v) in enumerate(zip(cells, rr, strict=True)):
                    c.text = str(v)
                    run = c.paragraphs[0].runs[0] if c.paragraphs[0].runs else None
                    if run is not None:
                        run.font.size = small
                        run.bold = b.kind == "kv" and j == 0
            doc.add_paragraph()
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def to_pdf(rep: Report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold = _pdf_fonts(pdfmetrics, TTFont)
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName=font, fontSize=8.5, leading=11)
    note = ParagraphStyle("n", parent=body, textColor=colors.HexColor("#555555"), fontName=font)
    cell = ParagraphStyle("c", parent=body, fontSize=7.5, leading=9)
    cellb = ParagraphStyle("cb", parent=cell, fontName=bold)
    caption = ParagraphStyle("cap", parent=body, fontName=bold, fontSize=8, spaceBefore=4, spaceAfter=2)
    heads = {
        "h0": ParagraphStyle("h0", parent=ss["Title"], fontName=bold, fontSize=16),
        "h1": ParagraphStyle("h1", parent=ss["Heading1"], fontName=bold, fontSize=13, spaceBefore=10),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontName=bold, fontSize=11),
        "h3": ParagraphStyle("h3", parent=ss["Heading3"], fontName=bold, fontSize=9.5),
    }
    esc = lambda t: str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731
    flow = [Paragraph(esc(rep.title), heads["h0"]), Paragraph(esc(rep.subtitle), body), Spacer(1, 4 * mm)]
    width = A4[0] - 30 * mm
    for b in rep.blocks:
        if b.kind.startswith("h"):
            flow.append(Paragraph(esc(b.text), heads[b.kind]))
        elif b.kind == "p":
            flow.append(Paragraph(esc(b.text), body))
        elif b.kind == "note":
            flow.append(Paragraph(esc(b.text), note))
        elif b.kind == "image":
            from reportlab.lib.utils import ImageReader
            from reportlab.platypus import Image as PdfImage

            iw, ih = ImageReader(io.BytesIO(b.image)).getSize()
            flow.append(PdfImage(io.BytesIO(b.image), width=width, height=width * ih / iw))
        elif b.kind == "caption":
            flow.append(Paragraph(esc(b.text), caption))
        elif b.kind == "bullets":
            flow += [Paragraph("• " + esc(x[0]), body) for x in b.rows]
        else:
            data = ([[Paragraph(esc(h), cellb) for h in b.headers]] if b.headers else []) + [
                [Paragraph(esc(v), cellb if b.kind == "kv" and j == 0 else cell) for j, v in enumerate(rr)]
                for rr in b.rows
            ]
            widths = [width * 0.3, width * 0.7] if b.kind == "kv" else [width * f for f in _shares(b)]
            t = Table(data, colWidths=widths, repeatRows=1 if b.headers else 0, hAlign="LEFT")
            style = [
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#bbbbbb")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
            if b.headers:
                style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef4")))
            t.setStyle(TableStyle(style))
            flow += [t, Spacer(1, 3 * mm)]
    out = io.BytesIO()

    def footer(canvas, doc):
        canvas.setFont(font, 7)
        canvas.drawString(15 * mm, 8 * mm, f"{rep.title} · {rep.subtitle}")
        canvas.drawRightString(A4[0] - 15 * mm, 8 * mm, f"Page {doc.page}")

    SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title=rep.title,
    ).build(flow, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


def _shares(b: Block) -> list[float]:
    """Column widths as shares of the page: by the longest text in each column, within limits."""
    table = ([b.headers] if b.headers else []) + b.rows
    longest = [max(len(str(v)) for v in col) for col in zip(*table, strict=False)]
    weights = [min(max(x, 6), 45) for x in longest]
    total = sum(weights)
    return [w / total for w in weights]


def _pdf_fonts(pdfmetrics, TTFont) -> tuple[str, str]:
    """A Unicode font for Ø, γ, σ and the like: DejaVu when the system has it, else Helvetica."""
    import os

    for folder in (
        "/usr/share/fonts/truetype/dejavu",
        "/usr/share/fonts/dejavu",
        "/Library/Fonts",
        "C:/Windows/Fonts",
    ):
        regular, bold = os.path.join(folder, "DejaVuSans.ttf"), os.path.join(folder, "DejaVuSans-Bold.ttf")
        if os.path.exists(regular) and os.path.exists(bold):
            if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("DejaVu", regular))
                pdfmetrics.registerFont(TTFont("DejaVu-Bold", bold))
            return "DejaVu", "DejaVu-Bold"
    for folder in ("C:/Windows/Fonts",):
        regular, bold = os.path.join(folder, "arial.ttf"), os.path.join(folder, "arialbd.ttf")
        if os.path.exists(regular):
            if "Arial" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("Arial", regular))
                pdfmetrics.registerFont(TTFont("Arial-Bold", bold))
            return "Arial", "Arial-Bold"
    return "Helvetica", "Helvetica-Bold"


RENDERERS = {
    "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "docx": (to_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pdf": (to_pdf, "application/pdf"),
}
