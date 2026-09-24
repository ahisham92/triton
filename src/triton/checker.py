"""The Checker workbook: the sheets that have warnings, as read (with any edits made in Triton),
flagged rows highlighted and the reason in a note, and a first sheet listing every warning with a
link to its rows.

Edited in Excel and uploaded again with "Replace matching tabs", its sheets replace the section's;
the list sheet is skipped on upload.
"""

from __future__ import annotations

import io
from collections.abc import Callable
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname

from .review import choices

LIST_SHEET = "Triton checker"
FILLS = {
    "error": PatternFill("solid", fgColor="F9D6D3"),
    "warning": PatternFill("solid", fgColor="FBE7C0"),
    "info": PatternFill("solid", fgColor="DCE8F4"),
}
_RANK = {"error": 0, "warning": 1, "info": 2}


def build(
    issues: list[dict[str, Any]], raw: Callable[[str], list | None], decisions: dict[str, str]
) -> bytes:
    """``issues`` as in the workbook summary; ``raw(sheet)`` gives a sheet's rows as read."""
    wb = Workbook()
    ws = wb.active
    ws.title = LIST_SHEET
    head = [
        "Severity",
        "Kind",
        "Sheet",
        "Element",
        "Combination",
        "Rows",
        "Warning",
        "Accept does",
        "Decision",
        "Go to",
    ]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    sheets: dict[str, dict[int, list[dict]]] = {}
    for i in issues:
        if i.get("sheet") and i.get("rows"):
            for r in i["rows"]:
                sheets.setdefault(i["sheet"], {}).setdefault(r, []).append(i)
        elif i.get("sheet"):
            sheets.setdefault(i["sheet"], {})
    for i in sorted(issues, key=lambda x: (_RANK.get(x["severity"], 3), x.get("sheet") or "")):
        rows = i.get("rows") or []
        c = choices(i["code"])
        ws.append(
            [
                i["severity"],
                i["code"].replace("_", " "),
                i.get("sheet") or "",
                i.get("element") or "",
                i.get("combination") or "",
                _ranges(rows),
                i["message"],
                c["accept"] if c else "Fix in the workbook, sheet mapping or load combinations",
                decisions.get(i["id"], "to review" if c and c["before"] != "auto" else ""),
                "",
            ]
        )
        row = ws.max_row
        ws.cell(row, 1).fill = FILLS.get(i["severity"], FILLS["info"])
        if i.get("sheet") and i["sheet"] in sheets:
            link = ws.cell(row, len(head))
            link.value = f"{i['sheet']}!A{rows[0] if rows else 1}"
            link.hyperlink = f"#{quote_sheetname(i['sheet'])}!A{rows[0] if rows else 1}"
            link.font = Font(color="1F5F8B", underline="single")
    for col, width in zip("ABCDEFGHIJ", (9, 22, 22, 14, 16, 16, 70, 40, 11, 24), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    for name, flags in sheets.items():
        rows = raw(name)
        out = wb.create_sheet(name[:31])
        if rows is None:
            out.append(
                [f"The rows of {name} were not kept (uploaded before the Checker existed): upload it again."]
            )
            continue
        for r in rows:
            out.append(list(r))
        for n, found in flags.items():
            worst = min(found, key=lambda x: _RANK.get(x["severity"], 3))
            fill = FILLS.get(worst["severity"], FILLS["info"])
            width = max(len(rows[n - 1]) if 0 < n <= len(rows) else 1, 1)
            for col in range(1, width + 1):
                out.cell(n, col).fill = fill
            note = "\n".join(dict.fromkeys(f"{x['severity'].upper()}: {x['message']}" for x in found))
            out.cell(n, 1).comment = Comment(note, "Triton", width=320, height=120)
        widest = max((len(r) for r in rows), default=0)
        for col in range(1, widest + 1):
            out.column_dimensions[get_column_letter(col)].width = 12
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _ranges(rows: list[int]) -> str:
    """12, 13, 14, 20 as "12-14, 20"."""
    out, start, prev = [], None, None
    for r in sorted(rows):
        if start is None:
            start = prev = r
        elif r == prev + 1:
            prev = r
        else:
            out.append(f"{start}" if start == prev else f"{start}-{prev}")
            start = prev = r
    if start is not None:
        out.append(f"{start}" if start == prev else f"{start}-{prev}")
    return ", ".join(out)
