"""Governing load sets per station, as the design office enters them in AdSec.

A station is a length of the element with one reinforcement cage (a run of the
curtailment, or the whole element when it is not reduced). For each station
and for ULS and QP separately there are seven sets:

    max N, min N     with the M2 and M3 at the same point
    max M2, min M2   with the N and M3 at the same point
    max M3, min M3   with the N and M2 at the same point
    most utilised    ULS: highest N–M utilisation with the station's cage
                     QP: largest resultant moment, until crack width is checked

The maxima and minima are taken over every combination of the element, not per
combination, and the most utilised row is kept even when it repeats one of the
six extremes. Points inside a steel casing have no crack width check, so they
are left out of the QP sets; where nothing is left (a station inside the
casing, the combi wall infill), the seven QP sets are 1s so the office's AdSec
template still has its 14 rows.

N follows the concrete (AdSec) sign convention: compression positive, i.e. the
Plaxis N multiplied by -1. M2 and M3 are the Plaxis values.

Steel elements (the combi wall tube, the sheet pile wall) are not designed with
these sets, so they get ten ULS rows instead: maximum and minimum N, M2, M3 and
the two shears, each with the other actions at the same point, in the Plaxis
sign.
"""

from __future__ import annotations

import math
from types import SimpleNamespace
from typing import Any

import numpy as np
import pandas as pd

from ..importer import SheetData
from ..project import DesignSettings, PileInput

# Plaxis columns of the five steel actions (N, M2, M3, Q1, Q2) for beam and plate results.
STEEL_COLUMNS = {
    "beam": ("N", "M_2", "M_3", "Q_12", "Q_13"),
    "plate": ("N_1", "M_11", "M_22", "Q_13", "Q_23"),
}
STEEL_KEYS = ("N", "M2", "M3", "Q1", "Q2")

EXTREMES = (
    ("max N", "N", True),
    ("min N", "N", False),
    ("max M2", "M_2", True),
    ("min M2", "M_2", False),
    ("max M3", "M_3", True),
    ("min M3", "M_3", False),
)


def _row(case: str, p: pd.Series, util: float | None) -> dict[str, Any]:
    return {
        "case": case,
        "combination": str(p["combination"]),
        "node": int(p["Node"]) if pd.notna(p.get("Node")) else None,
        "z": round(float(p["Z"]), 2),
        "N_kN": round(float(p["N"]), 1),
        "M2_kNm": round(float(p["M_2"]), 1),
        "M3_kNm": round(float(p["M_3"]), 1),
        "utilisation": None if util is None or not np.isfinite(util) else round(float(util), 3),
    }


def pick_sets(frame: pd.DataFrame, util: np.ndarray | None, seventh: str) -> list[dict[str, Any]]:
    """The six extremes and the ``seventh`` case of one station and limit state."""
    if frame.empty:
        return []
    frame = frame.reset_index(drop=True)
    rows = []
    for case, col, is_max in EXTREMES:
        i = int(frame[col].idxmax() if is_max else frame[col].idxmin())
        rows.append(_row(case, frame.iloc[i], None if util is None else util[i]))
    if util is not None:
        i = int(np.nanargmax(util))
    else:
        i = int(np.hypot(frame["M_2"], frame["M_3"]).idxmax())
    rows.append(_row(seventh, frame.iloc[i], None if util is None else util[i]))
    return rows


def placeholder_sets() -> list[dict[str, Any]]:
    """Seven QP rows of 1s, for stations without a crack width check."""
    cases = [c for c, _, _ in EXTREMES] + ["largest resultant M"]
    return [
        {
            "case": c,
            "combination": "no crack check",
            "node": None,
            "z": None,
            "N_kN": 1.0,
            "M2_kNm": 1.0,
            "M3_kNm": 1.0,
            "utilisation": None,
        }
        for c in cases
    ]


def cased(pile: PileInput, top: float, bottom: float) -> bool:
    """Whether a station lies wholly inside the pile's steel casing (no crack width check)."""
    c = pile.casing
    return c is not None and c.bottom_level <= bottom + 1e-9 and top <= c.top_level + 1e-9


def station_sets(
    pile: PileInput,
    settings: DesignSettings,
    uls: pd.DataFrame,
    qp: pd.DataFrame,
    stations: list[tuple[float, float, dict]],
) -> list[dict[str, Any]]:
    """Seven ULS and seven QP sets for each (top, bottom, cage dict) station."""
    from .piles import _utilisation

    out = []
    for top, bottom, cage in stations:
        arrangement = SimpleNamespace(rings=[SimpleNamespace(**r) for r in cage["rings"]])
        u_rows = uls[(uls["Z"] <= top + 1e-9) & (uls["Z"] >= bottom - 1e-9)]
        q_rows = qp[(qp["Z"] <= top + 1e-9) & (qp["Z"] >= bottom - 1e-9)] if not qp.empty else qp
        if pile.casing is not None and not q_rows.empty:
            c = pile.casing
            q_rows = q_rows[(q_rows["Z"] > c.top_level + 1e-9) | (q_rows["Z"] < c.bottom_level - 1e-9)]
        no_crack = cased(pile, top, bottom) or (pile.casing is not None and q_rows.empty)
        util = _utilisation(pile, arrangement, settings, u_rows) if len(u_rows) else None
        qp_sets = placeholder_sets()
        if not no_crack:
            qp_sets = pile_set_cracks(
                pile, settings, arrangement, pick_sets(q_rows, None, "largest resultant M")
            )
        out.append(
            {
                "top": top,
                "bottom": bottom,
                "cage": cage["label"],
                "rings": cage["rings"],
                "governing": _governing(pile, settings, arrangement, u_rows, util),
                "uls": pick_sets(u_rows, util, "most utilised"),
                "qp": qp_sets,
            }
        )
    return out


def pile_set_cracks(pile: PileInput, settings: DesignSettings, arrangement, rows: list[dict]) -> list[dict]:
    """Each QP set with its crack width and the terms that draw it (7.3.4 at the extreme bar)."""
    from .piles import crack_widths

    if not rows:
        return rows
    n = np.array([r["N_kN"] for r in rows], float)
    m2 = np.array([r["M2_kNm"] for r in rows], float)
    m3 = np.array([r["M3_kNm"] for r in rows], float)
    w = crack_widths(pile, arrangement, settings, pd.DataFrame({"N": n, "M": np.hypot(m2, m3)}))
    limit = pile.crack_width_limit
    for r, t, a2, a3 in zip(rows, w.itertuples(index=False), m2, m3, strict=True):
        r["crack"] = crack_terms(t.wk, limit, t.sigma_s, t.sr_max, t.x, pile.diameter)
        # The direction of the moment vector, from the M3 axis towards M2 (degrees).
        r["crack"]["angle_deg"] = round(math.degrees(math.atan2(a2, a3)), 1)
    return rows


def crack_terms(wk, limit, sigma_s, sr_max, x, h) -> dict:
    """What a crack picture needs: wk against its limit, the bar stress, the crack spacing and the
    compressed depth x (0 all in tension, h all in compression) of a section h deep (mm)."""

    def num(v, nd=0):
        return None if v is None or not np.isfinite(v) else round(float(v), nd) if nd else round(float(v))

    wk = float(wk or 0.0)
    return {
        "wk_mm": round(wk, 3),
        "limit_mm": limit,
        "util": round(wk / limit, 3) if limit else None,
        "sigma_s_MPa": num(sigma_s, 1),
        "sr_max_mm": num(sr_max) if wk > 0 else None,
        "x_mm": num(x),
        "h_mm": num(h),
    }


def _governing(
    pile: PileInput, settings: DesignSettings, arrangement, rows: pd.DataFrame, util
) -> dict | None:
    """The station's most utilised ULS point with MRd at its N, as in the design office summary tables."""
    from .piles import _accidental, _section

    if util is None or rows.empty:
        return None
    i = int(np.nanargmax(util))
    g = rows.iloc[i]
    sec = _section(pile, arrangement, settings, bool(_accidental(rows.iloc[[i]])[0]))
    n, m = float(g["N"]), float(np.hypot(g["M_2"], g["M_3"]))
    m_rd = sec.moment_capacity(n)
    return {
        "utilisation": round(float(util[i]), 3),
        "combination": str(g["combination"]),
        "z": round(float(g["Z"]), 2),
        "N_kN": round(n, 1),
        "M_kNm": round(m, 1),
        "M_Rd_kNm": round(m_rd, 1),
        "moment_ratio": round(m / m_rd, 3) if m_rd > 0 else None,
    }


def steel_sets(frame: pd.DataFrame, kind: str) -> dict[str, Any]:
    """Ten ULS rows of a steel element: max and min of N, M2, M3, Q1 and Q2 over all combinations.

    ``frame`` holds every ULS point with a ``combination`` column; ``kind`` is "beam" or "plate".
    """
    cols = STEEL_COLUMNS[kind]
    out = {"columns": dict(zip(STEEL_KEYS, cols, strict=True)), "rows": []}
    if frame.empty or not set(cols) <= set(frame.columns):
        return out
    frame = frame.reset_index(drop=True)
    for key, col in zip(STEEL_KEYS, cols, strict=True):
        for label, i in (("max", frame[col].idxmax()), ("min", frame[col].idxmin())):
            p = frame.loc[i]
            row = {
                "case": f"{label} {key}",
                "combination": str(p["combination"]),
                "node": int(p["Node"]) if pd.notna(p.get("Node")) else None,
                "z": round(float(p["Z"]), 2),
            }
            row.update({k: round(float(p[c]), 1) for k, c in zip(STEEL_KEYS, cols, strict=True)})
            out["rows"].append(row)
    return out


def steel_header(columns: dict[str, str]) -> list[str]:
    """Column titles of the steel rows, naming the Plaxis column where it differs (e.g. "Q1 (Q_12)")."""
    heads = [k if columns[k].replace("_", "") == k else f"{k} ({columns[k]})" for k in STEEL_KEYS]
    return ["Criterion", *heads, "Combination", "z (m)", "Node"]


def uls_frame(sheets: dict[str, SheetData]) -> pd.DataFrame:
    """Every ULS point of an element (QP sheets left out) with its combination."""
    from ..elements import CombinationType, combination_type

    parts = [
        s.frame.assign(combination=c)
        for c, s in sheets.items()
        if combination_type(c) is not CombinationType.SLS_QP and not s.frame.empty
    ]
    return pd.concat(parts, ignore_index=True) if parts else pd.DataFrame()


def qp_loads(sheets: dict[str, SheetData], head_level: float | None, above: float) -> pd.DataFrame:
    from .piles import PileLoads

    return PileLoads.from_sheets(sheets, head_level, above, qp=True).frame


CONCRETE_HEADER = [
    "Criterion",
    "N (kN)",
    "M2 (kNm)",
    "M3 (kNm)",
    "Combination",
    "z (m)",
    "Node",
    "Utilisation",
]


def workbook(project: str, section: str, results: dict[str, Any]) -> bytes:
    """Governing straining actions laid out for copy and paste into AdSec.

    ``Concrete``: for each pile, combi wall infill and beam (and each station where the cage
    changes down the element), its name, then 7 QP rows and 7 ULS rows underneath.
    ``Steel``: for each combi wall tube and sheet pile wall, its name and 10 ULS rows.
    """
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    bold = Font(bold=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Concrete"
    ws.append([f"{project} · {section} · concrete elements"])
    ws.append(
        ["N in the concrete (AdSec) sign convention: Plaxis N × −1, compression +. M2, M3 as in Plaxis."]
    )
    designs = [(p["element"], p) for p in results.get("piles", [])]
    designs += [(f"{w['element']} infill", w["infill"]) for w in results.get("combi_walls", [])]
    designs += [
        (
            f"{b['element']} · M3 vertical bending (sagging +), M2 horizontal; "
            "z = position along the beam (m)",
            b,
        )
        for b in results.get("beams", [])
    ]
    for name, d in designs:
        stations = d.get("governing_sets") or []
        for st in stations:
            ws.append([])
            title = name
            if len(stations) > 1 and "beam" not in d.get("kind", ""):
                title += f" · {st['top']:g} to {st['bottom']:g} m · {st['cage']}"
            ws.append([title])
            ws.cell(ws.max_row, 1).font = bold
            ws.append(CONCRETE_HEADER)
            for c in ws[ws.max_row]:
                c.font = bold
            for state, rows in (("QP", st["qp"]), ("ULS", st["uls"])):
                for r in rows:
                    ws.append(
                        [
                            f"{state} {r['case']}",
                            r["N_kN"],
                            r["M2_kNm"],
                            r["M3_kNm"],
                            r["combination"],
                            r["z"],
                            r["node"],
                            r["utilisation"],
                        ]
                    )
    _widths(ws, (26, 10, 10, 10, 16, 8, 9, 11))

    ss = wb.create_sheet("Steel")
    ss.append([f"{project} · {section} · steel elements"])
    ss.append(["Plaxis signs (N not multiplied by −1). Not designed in Triton: max and min of each action."])
    steel = [
        (f"{w['element']} tube", "kN, kNm", (w.get("tube") or {}).get("governing_sets"))
        for w in results.get("combi_walls", [])
    ]
    steel += [
        (s["element"], "kN/m, kNm/m", s.get("governing_sets")) for s in results.get("sheet_pile_walls", [])
    ]
    for name, unit, sets in steel:
        if not sets or not sets.get("rows"):
            continue
        cols = sets["columns"]
        ss.append([])
        ss.append([f"{name} ({unit})"])
        ss.cell(ss.max_row, 1).font = bold
        ss.append(steel_header(cols))
        for c in ss[ss.max_row]:
            c.font = bold
        for r in sets["rows"]:
            ss.append([f"ULS {r['case']}", *(r[k] for k in STEEL_KEYS), r["combination"], r["z"], r["node"]])
    _widths(ss, (22, 11, 11, 11, 11, 11, 16, 8, 9))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _widths(ws, widths: tuple[int, ...]) -> None:
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w
