"""Sheet pile wall straining actions for the design office's sheet pile program.

Triton does not design the sheet pile wall. It hands over the Plaxis plate
results, with the section's load multipliers applied, in the Plaxis sign
convention (steel element: N is not multiplied by -1):

* ``Governing``: for each combination, the maximum and minimum of every action
  with the other actions at the same node;
* one ``Envelope`` sheet per combination: maximum and minimum of every action
  across the wall at each level.
"""

from __future__ import annotations

import io
import re

import pandas as pd

from ..elements import PLATE_ACTIONS, combination_type
from ..importer import SheetData


def _actions(frame: pd.DataFrame) -> list[str]:
    return [a for a in PLATE_ACTIONS if a in frame.columns]


def governing_rows(combo: str, frame: pd.DataFrame) -> list[list]:
    acts = _actions(frame)
    rows = []
    for a in acts:
        for label, i in (("max", frame[a].idxmax()), ("min", frame[a].idxmin())):
            p = frame.loc[i]
            rows.append(
                [combo, f"{label} {a}", int(p["Node"]), round(float(p["Y"]), 2), round(float(p["Z"]), 2)]
                + [round(float(p[b]), 2) for b in acts]
            )
    return rows


def envelope(frame: pd.DataFrame) -> pd.DataFrame:
    acts = _actions(frame)
    g = frame.assign(level=frame["Z"].round(2)).groupby("level")[acts]
    out = pd.concat({"max": g.max(), "min": g.min()}, axis=1).sort_index(ascending=False)
    out.columns = [f"{a} {k}" for k, a in out.columns]
    return out[[f"{a} {k}" for a in acts for k in ("max", "min")]].round(2)


def workbook(project: str, section: str, element: str, sheets: dict[str, SheetData]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Governing"
    ws.append([f"{project} · {section} · {element}"])
    ws.append(["Plaxis signs (N not multiplied by -1), kN/m and kNm/m, load multipliers applied."])
    ws.append([])
    first = next(iter(sheets.values())).frame if sheets else pd.DataFrame(columns=PLATE_ACTIONS)
    ws.append(["Combination", "Case", "Node", "Y (m)", "Z (m)", *_actions(first)])
    for c in ws[4]:
        c.font = Font(bold=True)
    for combo, sheet in sheets.items():
        for row in governing_rows(combo, sheet.frame):
            ws.append(row)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 12
    for combo, sheet in sheets.items():
        env = envelope(sheet.frame)
        es = wb.create_sheet(_title(f"Envelope {combo}", wb.sheetnames))
        es.append([f"{element} · {combo} ({combination_type(combo).value}): max and min across the wall"])
        es.append([])
        es.append(["Z (m)", *env.columns])
        for c in es[3]:
            c.font = Font(bold=True)
        for z, vals in env.iterrows():
            es.append([float(z), *(float(v) for v in vals)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _title(name: str, taken: list[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "", name)[:31]
    out, i = base, 2
    while out in taken:
        out = f"{base[:28]} {i}"
        i += 1
    return out
