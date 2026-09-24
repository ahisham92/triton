"""Bar bending schedule (BS 8666:2020 style) of a designed section, as Excel.

Built from the same bars as the Revit export (``design.export.pile_cages``), so the schedule, the
drawings and the Revit cages agree. Each line is one bar mark: the element, where the bars sit, the
bar type and size (H = B500), the number per member and in total, the shape code and its dimensions,
the cut length and the weight.

How the lengths are taken:

* Straight bars (shape 00) at their designed length. Pile bars already include the anchorage above
  the head and the lap below each run. Bars longer than the longest standard cut length are
  scheduled as pieces of that length plus one closing piece, with a lap (Design settings, *Lap
  length*) at each joint. Bars stop at the cover from the member's ends and edges.
* Added slab bars run over their zone plus an anchorage of one lap length past each end, cut at
  the slab's edges.
* Pile links (shape 75): closed rings, L = π(A − d) + a lap, at the spacing of each link zone; inner
  rings of multi-row cages over the runs that have them.
* Beam links (shape 51): nested closed links, the outer one round the cage (A × B to the outside of
  the link) and inner ones round the inner legs; an odd leg is a single leg with 135° hooks (99).
  L = 2(A + B) + 20d, the BS 8666 allowance for two 135° hooks.
* Slab shear links (99): one leg the slab depth less the covers, with a hook round the top and the
  bottom bars, L = (h − 2c) + 24d.

Weights use d²/162.2 kg/m (BS 8666 Table 1). They come out a little above the design's steel
weights, which leave out laps, hooks and (for slabs) the shear links. The schedule is a draft for
the detailer: holes at the piles, the voids and bends at the edges are not deducted or added.
"""

from __future__ import annotations

import io
import math
import re
from collections import defaultdict
from typing import Any

from .design.export import named_parts, pile_cages
from .project import Project

HOOK_51 = 20  # × d: two 135° hooks on a closed link (BS 8666)
HOOK_99 = 24  # × d: a slab link's hooks round the top and bottom bars


def kg_per_m(d: float) -> float:
    return d * d / 162.2


def _prefix(name: str, taken: set[str]) -> str:
    """P1 for Pile(1), FB for Front Beam, D for Deck; made unique within the section."""
    words = re.findall(r"[A-Za-z]+|\d+", name)
    p = "".join(w[0].upper() if w.isalpha() else w for w in words) or "E"
    out, n = p, 2
    while out in taken:
        out, n = f"{p}{chr(64 + n)}", n + 1
    taken.add(out)
    return out


def _pieces(length_mm: float, d: float, stock_mm: float, lap_factor: float) -> list[tuple[float, int]]:
    """A run of ``length_mm`` as bars no longer than ``stock_mm``, lapped: [(length, number)]."""
    if length_mm <= stock_mm:
        return [(length_mm, 1)]
    lap = lap_factor * d
    n = math.ceil((length_mm - lap) / (stock_mm - lap))
    last = length_mm + (n - 1) * lap - (n - 1) * stock_mm
    return [(stock_mm, n - 1), (last, 1)]


class _Schedule:
    def __init__(self, bar_type: str) -> None:
        self.rows: list[dict[str, Any]] = []
        self.bar_type = bar_type
        self._marks: dict[str, int] = defaultdict(int)
        self._taken: set[str] = set()
        self._prefix: dict[str, str] = {}

    def add(
        self,
        element: str,
        where: str,
        d: float,
        per_member: int,
        members: int,
        shape: str,
        length_mm: float,
        dims: dict[str, float] | None = None,
        note: str = "",
    ) -> None:
        if per_member <= 0 or length_mm <= 0:
            return
        if element not in self._prefix:
            self._prefix[element] = _prefix(element, self._taken)
        self._marks[element] += 1
        length = 25 * math.ceil(length_mm / 25)  # cut lengths rounded up to 25 mm (BS 8666 8.3)
        total = per_member * members
        self.rows.append(
            {
                "mark": f"{self._prefix[element]}-{self._marks[element]:02d}",
                "element": element,
                "where": where,
                "type": f"{self.bar_type}{int(round(d))}",
                "d": int(round(d)),
                "per_member": per_member,
                "members": members,
                "total": total,
                "shape": shape,
                "dims": {k: round(v) for k, v in (dims or {"A": length}).items()},
                "length_mm": length,
                "kg": total * length / 1000 * kg_per_m(d),
                "note": note,
            }
        )


def _zones_by_element(results: dict[str, Any]) -> dict[str, list[dict]]:
    out = {p["element"]: (p.get("shear") or {}).get("zones") or [] for p in results.get("piles", [])}
    for w in results.get("combi_walls", []):
        out[w["infill"]["element"]] = (w["infill"].get("shear") or {}).get("zones") or []
    return out


def _link_count(top: float, bottom: float, spacing_mm: float) -> int:
    return max(0, math.floor((top - bottom) * 1000 / spacing_mm)) + 1


def _piles(
    s: _Schedule, data: dict, zones: dict[str, list[dict]], stock_mm: float, lap_factor: float
) -> None:
    for p in data["piles"]:
        name = p["element"] + (" infill" if p["part"] == "infill" else "")
        members = int(p.get("count") or 1)
        grouped: dict[tuple, int] = defaultdict(int)
        for r in p["runs"]:
            for row in r["rows"]:
                grouped[
                    (row["diameter_mm"], round(row["bar_length_m"] * 1000), r["top_m"], r["bottom_m"])
                ] += row["count"]
        for (d, length, top, bottom), n in grouped.items():
            for piece, k in _pieces(length, d, stock_mm, lap_factor):
                note = "lapped" if length > stock_mm else ""
                s.add(
                    name, f"Main bars, {top:+.2f} to {bottom:+.2f}", d, n * k, members, "00", piece, note=note
                )
        dl = float(p["link_diameter_mm"])
        a = p["diameter_mm"] - 2 * p["cover_mm"]  # outside of the ring
        lap = lap_factor * dl
        z = zones.get(p["element"]) or []
        if not z:
            continue
        for i, zone in enumerate(z):
            n = _link_count(zone["top"], zone["bottom"], zone["spacing_mm"]) - (1 if i else 0)
            s.add(
                name,
                f"Links {zone['top']:+.2f} to {zone['bottom']:+.2f} @ {zone['spacing_mm']:g}",
                dl,
                n,
                members,
                "75",
                math.pi * (a - dl) + lap,
                {"A": a, "B": lap},
            )
        spacing = min(zone["spacing_mm"] for zone in z)
        for r in p["runs"]:
            for hoop in r.get("inner_link_hoops_mm") or []:
                n = _link_count(r["top_m"], r["bottom_m"], spacing)
                s.add(
                    name,
                    f"Inner rings {r['top_m']:+.2f} to {r['bottom_m']:+.2f} @ {spacing:g}",
                    dl,
                    n,
                    members,
                    "75",
                    math.pi * (hoop - dl) + lap,
                    {"A": hoop, "B": lap},
                )


def _beams(s: _Schedule, data: dict, stock_mm: float, lap_factor: float) -> None:
    for b in data["beams"]:
        name = b["element"]
        c = b["cover_mm"]
        run = abs(b["end_m"] - b["start_m"]) * 1000 - 2 * c
        by: dict[tuple[str, float], int] = defaultdict(int)
        top = max(b["bars"], key=lambda x: x["z_mm"])
        bottom = min(b["bars"], key=lambda x: x["z_mm"])
        for bar in b["bars"]:
            # The top and bottom layers (a second layer sits about 2d deeper); the rest are side bars.
            z = bar["z_mm"]
            face = (
                "Top"
                if z >= top["z_mm"] - 3 * top["diameter_mm"]
                else "Bottom"
                if z <= bottom["z_mm"] + 3 * bottom["diameter_mm"]
                else "Side"
            )
            by[(face, bar["diameter_mm"])] += 1
        for (face, d), n in sorted(by.items(), key=lambda t: ("Top", "Bottom", "Side").index(t[0][0])):
            for length, k in _pieces(run, d, stock_mm, lap_factor):
                s.add(
                    name,
                    f"{face} bars",
                    d,
                    n * k,
                    1,
                    "00",
                    length,
                    note="lapped" if k and run > stock_mm else "",
                )
        links = b.get("links") or {}
        if links.get("spacing_mm"):
            dl, legs = float(links["diameter_mm"]), int(links.get("legs") or 2)
            a, bb = b["width_mm"] - 2 * c, b["depth_mm"] - 2 * c
            n = math.floor(run / links["spacing_mm"]) + 1
            pairs, gaps = legs // 2, max(legs - 1, 1)
            for k in range(pairs):
                ak = a * (gaps - 2 * k) / gaps
                if ak <= 0:
                    break
                s.add(
                    name,
                    f"Links @ {links['spacing_mm']:g}" + (" (outer)" if k == 0 else " (inner)"),
                    dl,
                    n,
                    1,
                    "51",
                    2 * (ak + bb) + HOOK_51 * dl,
                    {"A": ak, "B": bb},
                )
            if legs % 2:
                s.add(name, "Middle link leg", dl, n, 1, "99", bb + HOOK_51 * dl, {"A": bb}, "135° hooks")
        for face, t in (b.get("transverse") or {}).items():
            if not t or not t.get("spacing_mm"):
                continue
            n = (math.floor(run / t["spacing_mm"]) + 1) * int(t.get("layers") or 1)
            s.add(
                name,
                f"Transverse {face} @ {t['spacing_mm']:g}",
                t["diameter_mm"],
                n,
                1,
                "00",
                b["width_mm"] - 2 * c,
            )


def _slabs(s: _Schedule, data: dict, stock_mm: float, lap_factor: float) -> None:
    for sl in data["slabs"]:
        name = sl["element"]
        (x0, x1), (y0, y1) = sl["box_m"]["X"], sl["box_m"]["Y"]
        for f in sl["faces"]:
            c = f["cover_mm"]
            along_x = f["bars_along"] == "X"
            run = ((x1 - x0) if along_x else (y1 - y0)) * 1000 - 2 * c
            width = ((y1 - y0) if along_x else (x1 - x0)) * 1000 - 2 * c
            face = f"{f['face'].capitalize()} bars along {f['bars_along']}"
            for layer in (f.get("mesh") or {}).get("layers") or []:
                for bar in layer["bars"]:
                    if bar.get("kind") != "mesh":
                        continue
                    d, sp = bar["diameter_mm"], bar["spacing_mm"]
                    n = math.floor(width / sp) + 1
                    for length, k in _pieces(run, d, stock_mm, lap_factor):
                        s.add(name, f"{face}, mesh L{layer['layer']} @ {sp:g}", d, n * k, 1, "00", length)
            added: dict[tuple, int] = defaultdict(int)
            for z in f.get("zones") or []:
                (zx0, zx1), (zy0, zy1) = z["x_m"], z["y_m"]
                zl = ((zx1 - zx0) if along_x else (zy1 - zy0)) * 1000
                zw = ((zy1 - zy0) if along_x else (zx1 - zx0)) * 1000
                lo, hi = (zx0, zx1) if along_x else (zy0, zy1)
                b_lo, b_hi = (x0, x1) if along_x else (y0, y1)
                for layer in z.get("layers") or []:
                    for bar in layer["bars"]:
                        if bar.get("kind") == "mesh":
                            continue
                        d, sp = bar["diameter_mm"], bar["spacing_mm"]
                        anch = lap_factor * d
                        start = max(lo * 1000 - anch, b_lo * 1000 + c)
                        end = min(hi * 1000 + anch, b_hi * 1000 - c)
                        length = max(end - start, zl)
                        n = max(1, round(zw / sp))
                        added[(d, sp, layer["layer"], 25 * math.ceil(length / 25))] += n
            for (d, sp, lay, length), n in sorted(added.items()):
                for piece, k in _pieces(length, d, stock_mm, lap_factor):
                    s.add(
                        name, f"{face}, added L{lay} @ {sp:g}", d, n * k, 1, "00", piece, note="in the zones"
                    )
        h = sl["thickness_mm"]
        c = min((f["cover_mm"] for f in sl["faces"]), default=50)
        links: dict[float, int] = defaultdict(int)
        for lk in sl.get("links") or []:
            (lx0, lx1), (ly0, ly1) = lk["x_m"], lk["y_m"]
            nx = math.floor((lx1 - lx0) * 1000 / lk["sx_mm"]) + 1
            ny = math.floor((ly1 - ly0) * 1000 / lk["sy_mm"]) + 1
            links[lk["diameter_mm"]] += nx * ny
        for d, n in links.items():
            leg = h - 2 * c
            s.add(
                name,
                "Shear links",
                d,
                n,
                1,
                "99",
                leg + HOOK_99 * d,
                {"A": leg},
                "hooks round top and bottom bars",
            )


def schedule(project: Project, results: dict[str, Any], section: str = "") -> list[dict[str, Any]]:
    """Every bar mark of the designed piles, combi wall infills, beams and slabs in ``results``."""
    d = project.design
    grade = d.reinforcement.grade
    s = _Schedule("H" if grade.upper().startswith("B500") else grade)
    data = pile_cages(project.info.name, results, section=section)
    lap_factor = d.piles.lap_factor or 45.0
    stock_mm = max(d.piles.standard_bar_lengths or [12.0]) * 1000
    _piles(s, data, _zones_by_element(named_parts(results)), stock_mm, lap_factor)
    _beams(s, data, stock_mm, lap_factor)
    _slabs(s, data, stock_mm, lap_factor)
    return s.rows


def workbook(project: Project, section_name: str, rows: list[dict[str, Any]]) -> bytes:
    """The schedule as Excel: one sheet of bar marks, one of weights by element and bar size."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    bold = Font(bold=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    info = project.info
    ws.append([f"Bar bending schedule · {info.name} · {section_name}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(
        [
            f"Project number {info.number or '–'} · Grade {project.design.reinforcement.grade}"
            " · BS 8666:2020 shape codes"
        ]
    )
    ws.append([])
    head = [
        "Mark",
        "Element",
        "Where",
        "Type and size",
        "No. per member",
        "No. of members",
        "Total no.",
        "Length (mm)",
        "Shape code",
        "A (mm)",
        "B (mm)",
        "Weight (kg)",
        "Notes",
    ]
    ws.append(head)
    for c in ws[4]:
        c.font = bold
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        dims = r["dims"]
        ws.append(
            [
                r["mark"],
                r["element"],
                r["where"],
                r["type"],
                r["per_member"],
                r["members"],
                r["total"],
                r["length_mm"],
                r["shape"],
                dims.get("A"),
                dims.get("B"),
                round(r["kg"], 1),
                r["note"],
            ]
        )
    for col, w in zip("ABCDEFGHIJKLM", (9, 16, 34, 9, 9, 9, 9, 10, 8, 9, 9, 11, 28), strict=True):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"

    ws2 = wb.create_sheet("Weights")
    ws2.append(["Weight by element and bar size (t)"])
    ws2["A1"].font = Font(bold=True, size=13)
    sizes = sorted({r["d"] for r in rows})
    ws2.append(["Element", *[f"H{d}" for d in sizes], "Total"])
    for c in ws2[2]:
        c.font = bold
    elements = list(dict.fromkeys(r["element"] for r in rows))
    for e in elements:
        by = defaultdict(float)
        for r in rows:
            if r["element"] == e:
                by[r["d"]] += r["kg"] / 1000
        ws2.append([e, *[round(by.get(d, 0), 3) or None for d in sizes], round(sum(by.values()), 3)])
    tot = defaultdict(float)
    for r in rows:
        tot[r["d"]] += r["kg"] / 1000
    ws2.append(["Total", *[round(tot[d], 3) for d in sizes], round(sum(tot.values()), 3)])
    for c in ws2[ws2.max_row]:
        c.font = bold
    ws2.column_dimensions["A"].width = 20

    ws3 = wb.create_sheet("How lengths are taken")
    for line in (__doc__ or "").strip().splitlines():
        ws3.append([line])
    ws3.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
